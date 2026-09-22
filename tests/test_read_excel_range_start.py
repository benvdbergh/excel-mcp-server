"""read_excel_range keeps the caller start cell when end is omitted."""

from __future__ import annotations

import os
import sys

from openpyxl import Workbook

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from excel_mcp.data import read_excel_range  # noqa: E402


def test_read_excel_range_from_c5_excludes_a1(tmp_path) -> None:
    p = tmp_path / "data.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "origin"
    ws["C5"] = "target"
    wb.save(p)
    path = str(p.resolve())

    rows = read_excel_range(path, "S", "C5")
    flat = [cell for row in rows for cell in row]
    assert "target" in flat
    assert "origin" not in flat
