"""Chart data sheet and target_cell persist after save."""

from __future__ import annotations

import os
import sys

from openpyxl import Workbook, load_workbook

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from excel_mcp.chart import create_chart_in_sheet  # noqa: E402


def _book_with_data(tmp_path):
    p = tmp_path / "chart.xlsx"
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data["A1"] = "Cat"
    data["B1"] = "Val"
    data["A2"] = "X"
    data["B2"] = 1
    data["A3"] = "Y"
    data["B3"] = 2
    wb.create_sheet("Dashboard")
    wb.save(p)
    return str(p.resolve())


def test_chart_uses_qualified_data_sheet_and_anchor(tmp_path) -> None:
    path = _book_with_data(tmp_path)
    create_chart_in_sheet(
        path,
        "Dashboard",
        "Data!A1:B3",
        "bar",
        "E2",
        title="T",
    )
    ws = load_workbook(path)["Dashboard"]
    assert len(ws._charts) == 1
    chart = ws._charts[0]
    assert chart.anchor._from.col == 4
    assert chart.anchor._from.row == 1
    series_ref = chart.series[0].val.numRef.f
    assert "Data" in series_ref
    assert "Dashboard" not in series_ref


def test_chart_target_cell_aa10_does_not_raise(tmp_path) -> None:
    path = _book_with_data(tmp_path)
    create_chart_in_sheet(path, "Dashboard", "Data!A1:B3", "bar", "AA10")
    ws = load_workbook(path)["Dashboard"]
    chart = ws._charts[0]
    assert chart.anchor._from.col == 26
    assert chart.anchor._from.row == 9
