"""Tests for ``FileWorkbookService`` (STORY-3-1): delegation and contract surface."""

import json
import os
import sys
from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from excel_mcp.routing.file_workbook_service import FileWorkbookService  # noqa: E402
from excel_mcp.routing.routed_dispatch import (  # noqa: E402
    FILE_BACKEND_FORMULA_NOT_EVALUATED_CODE,
    file_backend_formula_not_evaluated_warning,
)
from excel_mcp.routing.workbook_operation_contract import (  # noqa: E402
    ROUTED_WORKBOOK_OPERATION_NAMES,
)


def test_file_workbook_service_has_all_routed_operation_names() -> None:
    svc = FileWorkbookService()
    for name in ROUTED_WORKBOOK_OPERATION_NAMES:
        assert hasattr(svc, name), name
        assert callable(getattr(svc, name)), name


def test_read_range_with_metadata_xlsm_formula_emits_warning(tmp_path) -> None:
    p = tmp_path / "macros.xlsm"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "=1+1"
    wb.save(p)
    path = str(p.resolve())

    warnings: list = []
    svc = FileWorkbookService()
    raw = svc.read_range_with_metadata(
        path,
        "Sheet1",
        "A1",
        "A1",
        operation_metadata={"_response_warnings": warnings},
    )
    data = json.loads(raw)
    assert data["cells"][0]["address"] == "A1"
    assert len(warnings) == 1
    assert warnings[0]["code"] == FILE_BACKEND_FORMULA_NOT_EVALUATED_CODE
    assert warnings[0] == file_backend_formula_not_evaluated_warning()


def test_read_range_with_metadata_xlsx_formula_no_warning(tmp_path) -> None:
    p = tmp_path / "plain.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "=1+1"
    wb.save(p)
    path = str(p.resolve())

    warnings: list = []
    svc = FileWorkbookService()
    svc.read_range_with_metadata(
        path,
        "Sheet1",
        "A1",
        "A1",
        operation_metadata={"_response_warnings": warnings},
    )
    assert warnings == []


@patch("excel_mcp.routing.file_workbook_service.read_excel_range_with_metadata")
def test_read_range_with_metadata_json(mock_read: MagicMock) -> None:
    payload = {
        "range": "A1:A1",
        "sheet_name": "S",
        "value_mode": "value",
        "cells": [{"address": "A1", "value": 1, "row": 1, "column": 1}],
    }
    mock_read.return_value = payload
    svc = FileWorkbookService()
    out = svc.read_range_with_metadata("/abs/book.xlsx", "S", "A1", "A1")
    mock_read.assert_called_once_with(
        "/abs/book.xlsx",
        "S",
        "A1",
        "A1",
        value_mode="value",
        metadata_mode="full",
        file_backend_warnings=None,
    )
    assert json.loads(out) == payload
    assert out.startswith("{")


@patch("excel_mcp.routing.file_workbook_service.read_excel_range_with_metadata")
def test_read_range_with_metadata_empty(mock_read: MagicMock) -> None:
    mock_read.return_value = {"cells": []}
    svc = FileWorkbookService()
    out = svc.read_range_with_metadata("/abs/b.xlsx", "S")
    data = json.loads(out)
    assert data["cells"] == []


def test_read_range_with_metadata_missing_sheet_returns_error(tmp_path) -> None:
    p = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.save(p)
    svc = FileWorkbookService()
    out = svc.read_range_with_metadata(str(p.resolve()), "Missing")
    assert out.startswith("Error:")
    assert "Missing" in out


def test_read_range_with_metadata_invalid_value_mode() -> None:
    svc = FileWorkbookService()
    with pytest.raises(ValueError, match="Invalid value_mode"):
        svc.read_range_with_metadata("/abs/b.xlsx", "S", value_mode="display")


def test_read_range_with_metadata_text_mode_file_backend(tmp_path) -> None:
    p = tmp_path / "formatted.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 19900
    ws["A1"].number_format = "#,##0.00"
    wb.save(p)
    path = str(p.resolve())

    svc = FileWorkbookService()
    raw = svc.read_range_with_metadata(path, "Sheet1", "A1", "A1", value_mode="text")
    data = json.loads(raw)
    assert data["value_mode"] == "text"
    assert data["cells"][0]["value"] == "19,900.00"
    assert data["cells"][0]["value"] != 19900


def test_read_range_with_metadata_default_echoes_value_mode(tmp_path) -> None:
    p = tmp_path / "plain.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 42
    wb.save(p)
    path = str(p.resolve())

    svc = FileWorkbookService()
    data = json.loads(svc.read_range_with_metadata(path, "Sheet1", "A1", "A1"))
    assert data["value_mode"] == "value"
    assert data["cells"][0]["value"] == 42


@patch("excel_mcp.routing.file_workbook_service.get_all_validation_ranges")
@patch("excel_mcp.routing.file_workbook_service.load_workbook")
def test_read_worksheet_data_validation_closes_workbook_on_get_all_failure(
    mock_load_workbook: MagicMock, mock_get_all: MagicMock
) -> None:
    ws = MagicMock()
    wb = MagicMock()
    wb.sheetnames = ["Data"]
    wb.__getitem__.return_value = ws
    mock_load_workbook.return_value = wb
    mock_get_all.side_effect = RuntimeError("validation scan failed")

    svc = FileWorkbookService()
    try:
        svc.read_worksheet_data_validation("/abs/w.xlsx", "Data")
    except RuntimeError:
        pass
    else:
        raise AssertionError("expected RuntimeError from get_all_validation_ranges")

    mock_load_workbook.assert_called_once_with("/abs/w.xlsx", read_only=False)
    wb.close.assert_called_once()


@patch("excel_mcp.routing.file_workbook_service.get_all_validation_ranges")
@patch("excel_mcp.routing.file_workbook_service.load_workbook")
def test_read_worksheet_data_validation(
    mock_load_workbook: MagicMock, mock_get_all: MagicMock
) -> None:
    ws = MagicMock()
    wb = MagicMock()
    wb.sheetnames = ["Data"]
    wb.__getitem__.return_value = ws
    mock_load_workbook.return_value = wb
    mock_get_all.return_value = [{"cells": "A1:A2", "type": "list"}]

    svc = FileWorkbookService()
    out = svc.read_worksheet_data_validation("/abs/w.xlsx", "Data")

    mock_load_workbook.assert_called_once_with("/abs/w.xlsx", read_only=False)
    mock_get_all.assert_called_once_with(ws)
    wb.close.assert_called_once()
    data = json.loads(out)
    assert data["sheet_name"] == "Data"
    assert data["validation_rules"] == mock_get_all.return_value


@patch("excel_mcp.routing.file_workbook_service.validate_range_in_sheet_operation")
def test_validate_sheet_range(mock_validate: MagicMock) -> None:
    mock_validate.return_value = {"message": "Range 'A1' is valid. Sheet contains data in range 'A1:B2'"}
    svc = FileWorkbookService()
    assert svc.validate_sheet_range("/abs/f.xlsx", "Sh", "A1") == mock_validate.return_value["message"]
    mock_validate.assert_called_once_with("/abs/f.xlsx", "Sh", "A1")

    svc.validate_sheet_range("/abs/f.xlsx", "Sh", "A1", "B2")
    mock_validate.assert_called_with("/abs/f.xlsx", "Sh", "A1:B2")


@patch("excel_mcp.routing.file_workbook_service.write_data")
def test_write_cell_grid(mock_write: MagicMock) -> None:
    mock_write.return_value = {"message": "Data written successfully to Sheet1"}
    svc = FileWorkbookService()
    grid = [[1, 2]]
    assert (
        svc.write_cell_grid("/abs/w.xlsx", "Sheet1", grid, "B2")
        == "Data written successfully to Sheet1"
    )
    mock_write.assert_called_once_with("/abs/w.xlsx", "Sheet1", grid, "B2")


@patch("excel_mcp.routing.file_workbook_service.wb_create_workbook")
def test_create_workbook(mock_create: MagicMock) -> None:
    svc = FileWorkbookService()
    out = svc.create_workbook("/abs/new.xlsx")
    mock_create.assert_called_once_with("/abs/new.xlsx")
    assert out == "Created workbook at /abs/new.xlsx"


@patch("excel_mcp.routing.file_workbook_service.load_workbook")
def test_save_workbook(mock_load_workbook: MagicMock) -> None:
    wb = MagicMock()
    mock_load_workbook.return_value = wb
    svc = FileWorkbookService()
    out = svc.save_workbook("/abs/existing.xlsx")
    mock_load_workbook.assert_called_once_with("/abs/existing.xlsx")
    wb.save.assert_called_once_with("/abs/existing.xlsx")
    wb.close.assert_called_once()
    assert out == "Workbook saved: /abs/existing.xlsx"


def test_routing_package_exports_file_workbook_service() -> None:
    from excel_mcp.routing import FileWorkbookService as FWS  # noqa: E402

    assert FWS is FileWorkbookService


def test_export_worksheet_table_file_backend(tmp_path) -> None:
    p = tmp_path / "table.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Col1", "Col2"])
    ws.append(["a", 1])
    ws.append(["b", 2])
    wb.save(p)
    path = str(p.resolve())

    svc = FileWorkbookService()
    data = json.loads(svc.export_worksheet_table(path, "Sheet1"))
    assert data["sheet_name"] == "Sheet1"
    assert data["headers"] == ["Col1", "Col2"]
    assert data["rows"] == [["a", 1], ["b", 2]]
    assert data["row_count"] == 2
    assert data["truncated"] is False
    assert data["max_rows"] == 10000


def test_export_worksheet_table_max_rows_truncates(tmp_path) -> None:
    p = tmp_path / "big.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["H"])
    for i in range(5):
        ws.append([i])
    wb.save(p)
    path = str(p.resolve())

    svc = FileWorkbookService()
    data = json.loads(svc.export_worksheet_table(path, "Sheet1", max_rows=2))
    assert data["headers"] == ["H"]
    assert len(data["rows"]) == 2
    assert data["row_count"] == 5
    assert data["truncated"] is True
    assert data["max_rows"] == 2


def test_export_worksheet_table_invalid_max_rows_file_backend(tmp_path) -> None:
    p = tmp_path / "big.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["H"])
    wb.save(p)
    path = str(p.resolve())

    svc = FileWorkbookService()
    out = svc.export_worksheet_table(path, "Sheet1", max_rows=0)
    assert out == "Error: max_rows must be a positive integer"


def test_list_tables_empty_workbook(tmp_path) -> None:
    p = tmp_path / "empty.xlsx"
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.save(p)
    path = str(p.resolve())
    mtime = p.stat().st_mtime_ns

    svc = FileWorkbookService()
    data = json.loads(svc.list_tables(path))
    assert data == {"tables": []}
    assert p.stat().st_mtime_ns == mtime


def test_list_tables_schema_and_minimal(tmp_path) -> None:
    from openpyxl.worksheet.table import Table

    p = tmp_path / "tables.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "software components"
    ws.append(["ColA", "ColB", "ColC"])
    ws.append([1, 2, 3])
    ws.append([4, 5, 6])
    ws.add_table(Table(displayName="ap_sw_components", ref="A1:C3"))
    plain = wb.create_sheet("stackfleet concept")
    plain.append(["X", "Y"])
    plain.append([9, 8])
    wb.save(p)
    path = str(p.resolve())
    mtime = p.stat().st_mtime_ns

    svc = FileWorkbookService()
    schema = json.loads(svc.list_tables(path, detail="schema"))
    assert len(schema["tables"]) == 1
    entry = schema["tables"][0]
    assert entry["sheet"] == "software components"
    assert entry["name"] == "ap_sw_components"
    assert entry["range"] == "A1:C3"
    assert entry["header_range"] == "A1:C1"
    assert entry["data_range"] == "A2:C3"
    assert entry["row_count"] == 2
    assert entry["filter_applied"] is False
    assert entry["columns"] == ["ColA", "ColB", "ColC"]
    # Catalog payload must not include cell values.
    assert set(entry) == {
        "sheet",
        "name",
        "range",
        "header_range",
        "data_range",
        "row_count",
        "filter_applied",
        "columns",
    }

    minimal = json.loads(svc.list_tables(path, detail="minimal"))
    assert "columns" not in minimal["tables"][0]
    assert minimal["tables"][0]["name"] == "ap_sw_components"
    assert p.stat().st_mtime_ns == mtime


def test_list_tables_filter_applied_from_autofilter(tmp_path) -> None:
    from openpyxl.worksheet.filters import AutoFilter, FilterColumn, Filters
    from openpyxl.worksheet.table import Table

    p = tmp_path / "filtered.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["H1", "H2"])
    ws.append([1, 2])
    ws.append([3, 4])
    table = Table(displayName="T1", ref="A1:B3")
    af = AutoFilter(ref="A1:B3")
    af.filterColumn.append(FilterColumn(colId=0, filters=Filters(filter=["1"])))
    table.autoFilter = af
    ws.add_table(table)
    wb.save(p)
    path = str(p.resolve())

    svc = FileWorkbookService()
    data = json.loads(svc.list_tables(path, detail="minimal"))
    assert data["tables"][0]["filter_applied"] is True


def test_list_tables_invalid_detail(tmp_path) -> None:
    p = tmp_path / "book.xlsx"
    wb = Workbook()
    wb.save(p)
    svc = FileWorkbookService()
    out = svc.list_tables(str(p.resolve()), detail="full")
    assert out.startswith("Error:")
    assert "detail" in out.lower()


def test_query_table_file_filter_and_no_save(tmp_path) -> None:
    from openpyxl.worksheet.table import Table

    p = tmp_path / "query.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Qty", "Tag"])
    ws.append(["a", 1, "x"])
    ws.append(["b", 2, "y"])
    ws.append(["c", 3, "x"])
    ws.add_table(Table(displayName="Items", ref="A1:C4"))
    wb.save(p)
    path = str(p.resolve())
    mtime = p.stat().st_mtime_ns

    svc = FileWorkbookService()
    data = json.loads(
        svc.query_table(
            path,
            "Items",
            columns=["Name", "Qty"],
            where=[{"column": "Tag", "op": "eq", "value": "x"}],
            limit=10,
            offset=0,
        )
    )
    assert data["headers"] == ["Name", "Qty"]
    assert data["rows"] == [{"Name": "a", "Qty": 1}, {"Name": "c", "Qty": 3}]
    assert data["row_count"] == 2
    assert data["truncated"] is False
    assert data["view_spec"]["target"]["name"] == "Items"
    assert data["view_applicability"]["limit_viewable"] is False
    assert data["view_applicability"]["viewable"] is True
    assert p.stat().st_mtime_ns == mtime


def test_query_table_file_unknown_table_and_column(tmp_path) -> None:
    from openpyxl.worksheet.table import Table

    p = tmp_path / "query_err.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["A", "B"])
    ws.append([1, 2])
    ws.add_table(Table(displayName="T1", ref="A1:B2"))
    wb.save(p)
    path = str(p.resolve())
    svc = FileWorkbookService()
    assert svc.query_table(path, "Missing").startswith("Error:")
    assert "Table" in svc.query_table(path, "Missing")
    bad_col = svc.query_table(path, "T1", columns=["Nope"])
    assert bad_col.startswith("Error:")
    assert "column" in bad_col.lower()


def test_map_sheet_layout_header_not_row1_and_excludes_table(tmp_path) -> None:
    from openpyxl.worksheet.table import Table

    p = tmp_path / "layout.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "stackfleet concept"
    # Blank rows 1–3; region header at row 4.
    ws["B4"] = "Name"
    ws["C4"] = "Qty"
    ws["B5"] = "alpha"
    ws["C5"] = 1
    ws["B6"] = "beta"
    ws["C6"] = 2
    # Native table elsewhere on the sheet (must not appear as an island).
    ws["E1"] = "TCol"
    ws["E2"] = 9
    ws.add_table(Table(displayName="SideTable", ref="E1:E2"))
    wb.save(p)
    path = str(p.resolve())
    before = p.read_bytes()

    svc = FileWorkbookService()
    data = json.loads(svc.map_sheet_layout(path, "stackfleet concept"))
    assert len(data["tables"]) == 1
    assert data["tables"][0]["kind"] == "table"
    assert data["tables"][0]["name"] == "SideTable"
    assert len(data["regions"]) == 1
    region = data["regions"][0]
    assert region["kind"] == "region"
    assert region["header_row"] == 4
    assert region["header_guess"] is True
    assert region["columns"] == ["Name", "Qty"]
    assert region["id"] == "stackfleet concept!B4:C6"
    assert region["range"] == "B4:C6"

    q = json.loads(
        svc.query_region(
            path,
            "stackfleet concept",
            region_id=region["id"],
            where=[{"column": "Name", "op": "eq", "value": "alpha"}],
        )
    )
    assert q["rows"] == [{"Name": "alpha", "Qty": 1}]
    assert q["view_spec"]["target"] == {
        "kind": "region",
        "sheet": "stackfleet concept",
        "range": "B4:C6",
    }
    assert p.read_bytes() == before


def test_query_region_explicit_range_and_errors(tmp_path) -> None:
    p = tmp_path / "region.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A3"] = "Name"
    ws["B3"] = "Tag"
    ws["A4"] = "a"
    ws["B4"] = "x"
    ws["A5"] = "b"
    ws["B5"] = "y"
    wb.save(p)
    path = str(p.resolve())
    mtime = p.stat().st_mtime_ns
    svc = FileWorkbookService()

    data = json.loads(
        svc.query_region(
            path,
            "Sheet1",
            range="A3:B5",
            columns=["Name"],
            where=[{"column": "Tag", "op": "eq", "value": "x"}],
        )
    )
    assert data["headers"] == ["Name"]
    assert data["rows"] == [{"Name": "a"}]
    assert p.stat().st_mtime_ns == mtime

    both = svc.query_region(
        path, "Sheet1", region_id="Sheet1!A3:B5", range="A3:B5"
    )
    assert both.startswith("Error:")
    assert "exactly one" in both.lower()
    unknown = svc.query_region(path, "Sheet1", region_id="Sheet1!Z9:Z10")
    assert unknown.startswith("Error:")
    assert "region" in unknown.lower()
