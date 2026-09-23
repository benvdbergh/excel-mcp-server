"""Numeric pivot row fields must still aggregate."""

from __future__ import annotations

import os
import sys

from openpyxl import Workbook, load_workbook

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from excel_mcp.fileio.pivot import create_pivot_table  # noqa: E402


def test_numeric_row_field_sums(tmp_path) -> None:
    p = tmp_path / "pivot.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Qty"
    ws["B1"] = "Amount"
    ws["A2"] = 5
    ws["B2"] = 10
    ws["A3"] = 5
    ws["B3"] = 3
    wb.save(p)
    path = str(p.resolve())

    create_pivot_table(path, "Data", "A1:B3", rows=["Qty"], values=["Amount"], agg_func="sum")
    out = load_workbook(path)
    pivot = out["Data_pivot"]
    amounts = [
        pivot.cell(row=r, column=2).value
        for r in range(2, pivot.max_row + 1)
        if str(pivot.cell(row=r, column=1).value) == "5"
    ]
    assert amounts == [13]
