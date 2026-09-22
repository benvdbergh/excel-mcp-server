"""delete_range_operation shifts only the selected rectangle."""

from __future__ import annotations

import os
import sys

from openpyxl import Workbook, load_workbook

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
_SRC = os.path.join(_REPO_ROOT, "src")
if _SRC not in sys.path:
    sys.path.insert(0, _SRC)

from excel_mcp.sheet import delete_range_operation  # noqa: E402


def test_shift_up_preserves_outside_columns(tmp_path) -> None:
    p = tmp_path / "shift.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A2"] = "A2"
    ws["B2"] = "B2"
    ws["C2"] = "C2"
    ws["D2"] = "D2"
    ws["B4"] = "B4"
    wb.save(p)
    path = str(p.resolve())

    delete_range_operation(path, "S", "B2", "C3", shift_direction="up")
    out = load_workbook(path)["S"]
    assert out["A2"].value == "A2"
    assert out["D2"].value == "D2"
    assert out["B2"].value == "B4"
    assert out["C2"].value is None


def test_shift_left_preserves_outside_rows(tmp_path) -> None:
    p = tmp_path / "shift-left.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "A1"
    ws["A2"] = "A2"
    ws["B2"] = "B2"
    ws["C2"] = "C2"
    ws["A3"] = "A3"
    wb.save(p)
    path = str(p.resolve())

    delete_range_operation(path, "S", "B2", "B2", shift_direction="left")
    out = load_workbook(path)["S"]
    assert out["A1"].value == "A1"
    assert out["A2"].value == "A2"
    assert out["A3"].value == "A3"
    assert out["B2"].value == "C2"
