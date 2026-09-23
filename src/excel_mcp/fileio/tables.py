import logging
import uuid
from typing import Any, Mapping, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

from excel_mcp.exceptions import DataError
from excel_mcp.query import (
    DEFAULT_LIST_TABLES_DETAIL,
    DETAIL_SCHEMA,
    build_sheet_layout_payload,
    build_table_catalog_entry,
    build_table_catalog_payload,
    layout_regions_from_grid,
    normalize_excel_a1_range,
    parse_region_id,
    query_region_rows,
    query_table_rows,
    resolve_region_query_source,
    validate_list_tables_detail,
)

logger = logging.getLogger(__name__)


def _table_ranges_from_ref(
    ref: str, *, header_row_count: int = 1, totals_row_count: int = 0
) -> tuple[str, Optional[str], int]:
    """Derive header range, data range, and data row count from a table ``ref``."""
    min_col, min_row, max_col, max_row = range_boundaries(normalize_excel_a1_range(ref))
    headers = max(0, int(header_row_count or 0))
    totals = max(0, int(totals_row_count or 0))
    if headers <= 0:
        headers = 1
    header_end = min_row + headers - 1
    header_range = (
        f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{header_end}"
    )
    data_start = header_end + 1
    data_end = max_row - totals
    if data_start > data_end:
        return header_range, None, 0
    data_range = (
        f"{get_column_letter(min_col)}{data_start}:"
        f"{get_column_letter(max_col)}{data_end}"
    )
    return header_range, data_range, data_end - data_start + 1


def _openpyxl_filter_applied(table: Table) -> bool:
    """True when the table AutoFilter stores at least one column filter criterion."""
    af = table.autoFilter
    if af is None:
        return False
    for col in af.filterColumn or ():
        if (
            col.filters is not None
            or col.customFilters is not None
            or col.colorFilter is not None
            or col.dynamicFilter is not None
            or col.iconFilter is not None
            or col.top10 is not None
        ):
            return True
    return False


def list_excel_tables(
    filepath: str, detail: str = DEFAULT_LIST_TABLES_DETAIL
) -> dict[str, Any]:
    """List native Excel tables (openpyxl) without writing the workbook."""
    mode = validate_list_tables_detail(detail)
    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        tables: list[dict[str, Any]] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for table_name in list(ws.tables):
                table = ws.tables[table_name]
                ref = str(table.ref)
                header_range, data_range, row_count = _table_ranges_from_ref(
                    ref,
                    header_row_count=int(table.headerRowCount or 1),
                    totals_row_count=int(table.totalsRowCount or 0),
                )
                columns: Optional[list[str]] = None
                if mode == DETAIL_SCHEMA:
                    columns = list(table.column_names)
                tables.append(
                    build_table_catalog_entry(
                        sheet=sheet_name,
                        name=str(table.displayName or table.name or table_name),
                        range_=ref,
                        header_range=header_range,
                        data_range=data_range,
                        row_count=row_count,
                        filter_applied=_openpyxl_filter_applied(table),
                        columns=columns,
                    )
                )
        return build_table_catalog_payload(tables)
    finally:
        wb.close()


def _find_openpyxl_table(
    wb: Any, table_name: str
) -> tuple[Any, Any, str]:
    """Return ``(worksheet, table, display_name)`` or raise ``DataError``."""
    target = table_name.strip()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for key in list(ws.tables):
            table = ws.tables[key]
            display = str(table.displayName or table.name or key)
            if display == target or str(key) == target:
                return ws, table, display
    raise DataError(f"Table '{table_name}' not found")


def query_excel_table(
    filepath: str,
    table: str,
    columns: Optional[Sequence[str]] = None,
    where: Optional[Sequence[Mapping[str, Any]]] = None,
    limit: Optional[int] = None,
    offset: Optional[int] = None,
    omit_empty: bool = False,
    search: Optional[str] = None,
) -> dict[str, Any]:
    """Query a native Excel table via openpyxl without saving the workbook."""
    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        ws, tbl, display_name = _find_openpyxl_table(wb, table)
        table_columns = list(tbl.column_names)
        min_col, min_row, max_col, max_row = range_boundaries(
            normalize_excel_a1_range(str(tbl.ref))
        )
        headers_count = int(tbl.headerRowCount or 1) or 1
        totals = int(tbl.totalsRowCount or 0)
        data_start = min_row + headers_count
        data_end = max_row - totals
        rows: list[dict[str, Any]] = []
        if data_start <= data_end:
            for r in range(data_start, data_end + 1):
                row_map: dict[str, Any] = {}
                for i, col_name in enumerate(table_columns):
                    c = min_col + i
                    row_map[col_name] = ws.cell(row=r, column=c).value
                rows.append(row_map)
        return query_table_rows(
            table_name=display_name,
            table_columns=table_columns,
            rows=rows,
            columns=columns,
            where=where,
            limit=limit,
            offset=offset,
            omit_empty=omit_empty,
            search=search,
        )
    finally:
        wb.close()


def _is_empty_cell(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def _openpyxl_sheet_content_grid(
    ws: Any,
) -> tuple[Optional[tuple[int, int, int, int]], list[list[Any]]]:
    """Return content bounds and value grid, or ``(None, [])`` when empty."""
    max_row = int(ws.max_row or 1)
    max_col = int(ws.max_column or 1)
    min_r = min_c = max_r = max_c = None
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            if not _is_empty_cell(ws.cell(row=r, column=c).value):
                if min_r is None:
                    min_r = max_r = r
                    min_c = max_c = c
                else:
                    min_r = min(min_r, r)
                    max_r = max(max_r, r)
                    min_c = min(min_c, c)
                    max_c = max(max_c, c)
    if min_r is None:
        return None, []
    assert min_c is not None and max_r is not None and max_c is not None
    grid: list[list[Any]] = []
    for r in range(min_r, max_r + 1):
        grid.append(
            [ws.cell(row=r, column=c).value for c in range(min_c, max_c + 1)]
        )
    return (min_r, min_c, max_r, max_c), grid


def _openpyxl_sheet_tables(
    ws: Any, sheet_name: str, *, detail: str = DETAIL_SCHEMA
) -> list[dict[str, Any]]:
    """Catalog ListObjects on one sheet (layout ``kind: table`` entries)."""
    mode = validate_list_tables_detail(detail)
    tables: list[dict[str, Any]] = []
    for table_name in list(ws.tables):
        table = ws.tables[table_name]
        ref = str(table.ref)
        header_range, data_range, row_count = _table_ranges_from_ref(
            ref,
            header_row_count=int(table.headerRowCount or 1),
            totals_row_count=int(table.totalsRowCount or 0),
        )
        columns: Optional[list[str]] = None
        if mode == DETAIL_SCHEMA:
            columns = list(table.column_names)
        entry = build_table_catalog_entry(
            sheet=sheet_name,
            name=str(table.displayName or table.name or table_name),
            range_=ref,
            header_range=header_range,
            data_range=data_range,
            row_count=row_count,
            filter_applied=_openpyxl_filter_applied(table),
            columns=columns,
        )
        entry["kind"] = "table"
        tables.append(entry)
    return tables


def map_sheet_layout_from_worksheet(ws: Any, sheet_name: str) -> dict[str, Any]:
    """Map tables and islands for an already-open worksheet."""
    tables = _openpyxl_sheet_tables(ws, sheet_name)
    excluded = [str(t["range"]) for t in tables]
    bounds, grid = _openpyxl_sheet_content_grid(ws)
    if bounds is None:
        regions: list[dict[str, Any]] = []
    else:
        min_r, min_c, _, _ = bounds
        regions = layout_regions_from_grid(
            sheet=sheet_name,
            values=grid,
            origin_row=min_r,
            origin_col=min_c,
            excluded_ranges=excluded,
        )
    return build_sheet_layout_payload(tables, regions)


def map_excel_sheet_layout(filepath: str, sheet_name: str) -> dict[str, Any]:
    """Map native tables and non-table islands on one sheet (openpyxl, no save)."""
    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise DataError(f"Sheet '{sheet_name}' not found")
        return map_sheet_layout_from_worksheet(wb[sheet_name], sheet_name)
    finally:
        wb.close()


def _read_openpyxl_range_matrix(ws: Any, range_a1: str) -> list[list[Any]]:
    min_col, min_row, max_col, max_row = range_boundaries(
        normalize_excel_a1_range(range_a1)
    )
    matrix: list[list[Any]] = []
    for r in range(min_row, max_row + 1):
        matrix.append(
            [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
        )
    return matrix


def query_excel_region(
    filepath: str,
    sheet_name: str,
    *,
    region_id: Optional[str] = None,
    range: Optional[str] = None,  # noqa: A002 — MCP arg name
    columns: Optional[Sequence[str]] = None,
    where: Optional[Sequence[Mapping[str, Any]]] = None,
    limit: Optional[int] = None,
    offset: Optional[int] = None,
    omit_empty: bool = False,
    search: Optional[str] = None,
) -> dict[str, Any]:
    """Query a non-table region by id or explicit A1 range (openpyxl, no save)."""
    has_id = region_id is not None and str(region_id).strip() != ""
    has_range = range is not None and str(range).strip() != ""
    if has_id == has_range:
        raise DataError("Provide exactly one of region_id or range")

    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise DataError(f"Sheet '{sheet_name}' not found")
        ws = wb[sheet_name]
        if has_id:
            range_a1 = parse_region_id(str(region_id), sheet_name=sheet_name)
            layout = map_sheet_layout_from_worksheet(ws, sheet_name)
            known_ids = {str(r["id"]) for r in layout["regions"]}
            if str(region_id).strip() not in known_ids:
                raise DataError(f"Unknown region_id: {region_id!r}")
            require_guess = True
        else:
            range_a1 = normalize_excel_a1_range(str(range))
            require_guess = False

        matrix = _read_openpyxl_range_matrix(ws, range_a1)
        region_columns, rows, rng = resolve_region_query_source(
            matrix=matrix,
            range_a1=range_a1,
            require_header_guess=require_guess,
        )
        return query_region_rows(
            sheet=sheet_name,
            range_a1=rng,
            region_columns=region_columns,
            rows=rows,
            columns=columns,
            where=where,
            limit=limit,
            offset=offset,
            omit_empty=omit_empty,
            search=search,
        )
    finally:
        wb.close()


def create_excel_table(
    filepath: str,
    sheet_name: str,
    data_range: str,
    table_name: str | None = None,
    table_style: str = "TableStyleMedium9"
) -> dict:
    """Creates a native Excel table for the given data range.
    
    Args:
        filepath: Path to the Excel file.
        sheet_name: Name of the worksheet.
        data_range: The cell range for the table (e.g., "A1:D5").
        table_name: A unique name for the table. If not provided, a unique name is generated.
        table_style: The visual style to apply to the table.
        
    Returns:
        A dictionary with a success message and table details.
    """
    try:
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise DataError(f"Sheet '{sheet_name}' not found.")
            
        ws = wb[sheet_name]

        # If no table name is provided, generate a unique one
        if not table_name:
            table_name = f"Table_{uuid.uuid4().hex[:8]}"

        # Check if table name already exists
        if table_name in ws.parent.defined_names:
            raise DataError(f"Table name '{table_name}' already exists.")

        # Create the table
        table = Table(displayName=table_name, ref=data_range)
        
        # Apply style
        style = TableStyleInfo(
            name=table_style, 
            showFirstColumn=False,
            showLastColumn=False, 
            showRowStripes=True, 
            showColumnStripes=False
        )
        table.tableStyleInfo = style
        
        ws.add_table(table)
        
        wb.save(filepath)
        
        return {
            "message": f"Successfully created table '{table_name}' in sheet '{sheet_name}'.",
            "table_name": table_name,
            "range": data_range
        }

    except Exception as e:
        logger.error(f"Failed to create table: {e}")
        raise DataError(str(e))
