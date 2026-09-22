import logging
import uuid
from typing import Any, Mapping, Optional, Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

from .exceptions import DataError

logger = logging.getLogger(__name__)

DETAIL_MINIMAL = "minimal"
DETAIL_SCHEMA = "schema"
DEFAULT_LIST_TABLES_DETAIL = DETAIL_SCHEMA
_VALID_LIST_TABLES_DETAIL = frozenset({DETAIL_MINIMAL, DETAIL_SCHEMA})

# query_table defaults (documented in TOOLS.md)
DEFAULT_QUERY_TABLE_LIMIT = 100
MAX_QUERY_TABLE_COLUMNS_WITHOUT_EXPLICIT = 32

_VIEWABLE_OPS = frozenset(
    {"eq", "neq", "contains", "in", "gt", "gte", "lt", "lte"}
)
_ALL_QUERY_OPS = _VIEWABLE_OPS | {"is_empty"}
_OPS_REQUIRING_VALUE = frozenset(
    {"eq", "neq", "contains", "in", "gt", "gte", "lt", "lte"}
)

# apply_table_view / clear_table_view (Epic 14 / STORY-14-1)
VIEW_MODE_IN_PLACE = "in_place"
VIEW_MODE_SNAPSHOT = "snapshot"
_VALID_VIEW_MODES = frozenset({VIEW_MODE_IN_PLACE, VIEW_MODE_SNAPSHOT})
_VALID_SORT_ORDERS = frozenset({"asc", "desc"})
RESTORE_TOKEN_VERSION = 1
# Excel XlAutoFilterOperator (numeric; avoid win32com import in shared module)
XL_FILTER_OR = 2
XL_FILTER_VALUES = 7
XL_SORT_ASCENDING = 1
XL_SORT_DESCENDING = 2


def normalize_excel_a1_range(addr: str) -> str:
    """Strip sheet qualifiers and ``$`` from an Excel A1 range address."""
    s = str(addr).strip()
    if "!" in s:
        s = s.split("!", 1)[1]
    return s.replace("$", "")


def validate_list_tables_detail(detail: str) -> str:
    """Return normalized detail mode or raise ``DataError``."""
    mode = (detail or DEFAULT_LIST_TABLES_DETAIL).strip().lower()
    if mode not in _VALID_LIST_TABLES_DETAIL:
        raise DataError(
            f"Invalid detail: {detail!r}; expected '{DETAIL_MINIMAL}' or '{DETAIL_SCHEMA}'"
        )
    return mode


def build_table_catalog_entry(
    *,
    sheet: str,
    name: str,
    range_: str,
    header_range: str,
    data_range: Optional[str],
    row_count: int,
    filter_applied: bool,
    columns: Optional[list[str]] = None,
) -> dict[str, Any]:
    """One ListObject catalog row (no cell values)."""
    entry: dict[str, Any] = {
        "sheet": sheet,
        "name": name,
        "range": normalize_excel_a1_range(range_),
        "header_range": normalize_excel_a1_range(header_range),
        "data_range": (
            normalize_excel_a1_range(data_range) if data_range else None
        ),
        "row_count": int(row_count),
        "filter_applied": bool(filter_applied),
    }
    if columns is not None:
        entry["columns"] = list(columns)
    return entry


def build_table_catalog_payload(tables: list[Mapping[str, Any]]) -> dict[str, Any]:
    """Stable workbook-scoped catalog envelope."""
    return {"tables": list(tables)}


def _table_ranges_from_ref(
    ref: str, *, header_row_count: int = 1, totals_row_count: int = 0
) -> tuple[str, Optional[str], int]:
    """Derive header range, data range, and data row count from a table ``ref``."""
    min_col, min_row, max_col, max_row = range_boundaries(normalize_excel_a1_range(ref))
    headers = max(0, int(header_row_count or 0))
    totals = max(0, int(totals_row_count or 0))
    if headers <= 0:
        headers = 1
    header_end = min_row + headers - 1
    header_range = (
        f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{header_end}"
    )
    data_start = header_end + 1
    data_end = max_row - totals
    if data_start > data_end:
        return header_range, None, 0
    data_range = (
        f"{get_column_letter(min_col)}{data_start}:"
        f"{get_column_letter(max_col)}{data_end}"
    )
    return header_range, data_range, data_end - data_start + 1


def _openpyxl_filter_applied(table: Table) -> bool:
    """True when the table AutoFilter stores at least one column filter criterion."""
    af = table.autoFilter
    if af is None:
        return False
    for col in af.filterColumn or ():
        if (
            col.filters is not None
            or col.customFilters is not None
            or col.colorFilter is not None
            or col.dynamicFilter is not None
            or col.iconFilter is not None
            or col.top10 is not None
        ):
            return True
    return False


def list_excel_tables(
    filepath: str, detail: str = DEFAULT_LIST_TABLES_DETAIL
) -> dict[str, Any]:
    """List native Excel tables (openpyxl) without writing the workbook."""
    mode = validate_list_tables_detail(detail)
    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        tables: list[dict[str, Any]] = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for table_name in list(ws.tables):
                table = ws.tables[table_name]
                ref = str(table.ref)
                header_range, data_range, row_count = _table_ranges_from_ref(
                    ref,
                    header_row_count=int(table.headerRowCount or 1),
                    totals_row_count=int(table.totalsRowCount or 0),
                )
                columns: Optional[list[str]] = None
                if mode == DETAIL_SCHEMA:
                    columns = list(table.column_names)
                tables.append(
                    build_table_catalog_entry(
                        sheet=sheet_name,
                        name=str(table.displayName or table.name or table_name),
                        range_=ref,
                        header_range=header_range,
                        data_range=data_range,
                        row_count=row_count,
                        filter_applied=_openpyxl_filter_applied(table),
                        columns=columns,
                    )
                )
        return build_table_catalog_payload(tables)
    finally:
        wb.close()


def _try_number(value: Any) -> Optional[float]:
    """Parse a numeric value for comparison; ``bool`` is not treated as numeric."""
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    return None


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _is_empty_cell(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str) and value.strip() == "":
        return True
    return False


def _values_equal(left: Any, right: Any) -> bool:
    ln = _try_number(left)
    rn = _try_number(right)
    if ln is not None and rn is not None:
        return ln == rn
    return _as_text(left) == _as_text(right)


def _compare_ordered(left: Any, right: Any, op: str) -> bool:
    ln = _try_number(left)
    rn = _try_number(right)
    if ln is not None and rn is not None:
        a, b = ln, rn
    else:
        a, b = _as_text(left), _as_text(right)
    if op == "gt":
        return a > b
    if op == "gte":
        return a >= b
    if op == "lt":
        return a < b
    if op == "lte":
        return a <= b
    raise DataError(f"Unsupported comparison operator: {op!r}")


def evaluate_row_clause(row: Mapping[str, Any], clause: Mapping[str, Any]) -> bool:
    """Apply one ``where`` clause to a row dict keyed by column header."""
    column = clause["column"]
    op = clause["op"]
    cell = row.get(column)
    if op == "is_empty":
        return _is_empty_cell(cell)
    value = clause.get("value")
    if op == "eq":
        return _values_equal(cell, value)
    if op == "neq":
        return not _values_equal(cell, value)
    if op == "contains":
        return _as_text(value).lower() in _as_text(cell).lower()
    if op == "in":
        if not isinstance(value, (list, tuple)):
            raise DataError("Operator 'in' requires a list value")
        return any(_values_equal(cell, item) for item in value)
    if op in ("gt", "gte", "lt", "lte"):
        return _compare_ordered(cell, value, op)
    raise DataError(f"Unsupported where operator: {op!r}")


def normalize_where_clauses(
    where: Optional[Sequence[Mapping[str, Any]]],
    *,
    known_columns: Sequence[str],
) -> list[dict[str, Any]]:
    """Validate and normalize ``where`` clauses; raise ``DataError`` on bad input."""
    known = set(known_columns)
    if where is None:
        return []
    if not isinstance(where, (list, tuple)):
        raise DataError("where must be a list of clause objects")
    out: list[dict[str, Any]] = []
    for i, raw in enumerate(where):
        if not isinstance(raw, Mapping):
            raise DataError(f"Invalid where clause at index {i}: expected object")
        column = raw.get("column")
        op_raw = raw.get("op")
        if not isinstance(column, str) or not column:
            raise DataError(f"Invalid where clause at index {i}: column is required")
        if column not in known:
            raise DataError(f"Unknown column: {column!r}")
        if not isinstance(op_raw, str) or not op_raw:
            raise DataError(f"Invalid where clause at index {i}: op is required")
        op = op_raw.strip().lower()
        if op not in _ALL_QUERY_OPS:
            raise DataError(
                f"Invalid where operator: {op_raw!r}; expected one of "
                f"{', '.join(sorted(_ALL_QUERY_OPS))}"
            )
        clause: dict[str, Any] = {"column": column, "op": op}
        if op in _OPS_REQUIRING_VALUE:
            if "value" not in raw:
                raise DataError(
                    f"Invalid where clause at index {i}: value is required for op {op!r}"
                )
            value = raw["value"]
            if op == "in" and not isinstance(value, (list, tuple)):
                raise DataError(
                    f"Invalid where clause at index {i}: 'in' value must be a list"
                )
            clause["value"] = list(value) if op == "in" else value
        out.append(clause)
    return out


def resolve_query_columns(
    table_columns: Sequence[str],
    columns: Optional[Sequence[str]],
) -> list[str]:
    """Resolve projection columns; enforce width threshold when omitted."""
    names = list(table_columns)
    if columns is None:
        if len(names) > MAX_QUERY_TABLE_COLUMNS_WITHOUT_EXPLICIT:
            raise DataError(
                f"Table has {len(names)} columns; pass an explicit columns list "
                f"(threshold is {MAX_QUERY_TABLE_COLUMNS_WITHOUT_EXPLICIT})"
            )
        return names
    if not isinstance(columns, (list, tuple)):
        raise DataError("columns must be a list of column names")
    known = set(names)
    out: list[str] = []
    for name in columns:
        if not isinstance(name, str) or not name:
            raise DataError(f"Invalid column name: {name!r}")
        if name not in known:
            raise DataError(f"Unknown column: {name!r}")
        if name not in out:
            out.append(name)
    return out


def normalize_query_pagination(
    limit: Optional[int], offset: Optional[int]
) -> tuple[int, int]:
    """Apply default limit and validate offset/limit."""
    if offset is None:
        off = 0
    else:
        try:
            off = int(offset)
        except (TypeError, ValueError) as e:
            raise DataError(f"Invalid offset: {offset!r}") from e
        if off < 0:
            raise DataError(f"Invalid offset: {offset!r}; must be >= 0")
    if limit is None:
        lim = DEFAULT_QUERY_TABLE_LIMIT
    else:
        try:
            lim = int(limit)
        except (TypeError, ValueError) as e:
            raise DataError(f"Invalid limit: {limit!r}") from e
        if lim < 0:
            raise DataError(f"Invalid limit: {limit!r}; must be >= 0")
    return lim, off


def build_view_applicability(
    clauses: Sequence[Mapping[str, Any]],
    *,
    offset: int,
    truncated: bool,
) -> dict[str, Any]:
    """Describe which parts of the query Excel can later show as a view."""
    clause_apps: list[dict[str, Any]] = []
    for clause in clauses:
        op = str(clause["op"])
        viewable = op in _VIEWABLE_OPS
        entry: dict[str, Any] = {
            "column": clause["column"],
            "op": op,
            "viewable": viewable,
        }
        if not viewable:
            entry["reason"] = (
                f"Operator {op!r} is not expressible as an Excel AutoFilter view"
            )
        clause_apps.append(entry)

    all_clauses_ok = all(c["viewable"] for c in clause_apps)
    page_complete = offset == 0 and not truncated
    overall = all_clauses_ok and page_complete
    result: dict[str, Any] = {
        "viewable": overall,
        "limit_viewable": False,
        "offset_viewable": False,
        "clauses": clause_apps,
    }
    if all_clauses_ok and not page_complete:
        reasons: list[str] = []
        if offset > 0:
            reasons.append("offset drops leading matching rows")
        if truncated:
            reasons.append("limit truncates matching rows")
        result["reason"] = "; ".join(reasons) or "limit/offset changed the result set"
    return result


def normalize_view_sort(
    sort: Any,
    *,
    known_columns: Sequence[str],
) -> Optional[dict[str, Any]]:
    """Normalize ``view_spec.sort`` to ``{\"by\": [{\"column\", \"order\"}]}`` or ``None``.

    Accepted shapes:
    - ``None``
    - ``{\"by\": [{\"column\": \"Qty\", \"order\": \"asc\"|\"desc\"}, ...]}``
    - ``{\"column\": \"Qty\", \"order\": \"asc\"|\"desc\"}`` (single-key sugar)
    """
    if sort is None:
        return None
    if not isinstance(sort, Mapping):
        raise DataError("view_spec.sort must be an object or null")
    known = set(known_columns)
    if "by" in sort:
        raw_by = sort["by"]
        if not isinstance(raw_by, (list, tuple)) or not raw_by:
            raise DataError("view_spec.sort.by must be a non-empty list")
        entries = list(raw_by)
    elif "column" in sort:
        entries = [sort]
    else:
        raise DataError(
            "view_spec.sort must include 'by' or 'column'; "
            "expected {\"by\": [{\"column\", \"order\"}]}"
        )
    out: list[dict[str, str]] = []
    for i, entry in enumerate(entries):
        if not isinstance(entry, Mapping):
            raise DataError(f"Invalid sort entry at index {i}: expected object")
        column = entry.get("column")
        order_raw = entry.get("order", "asc")
        if not isinstance(column, str) or not column:
            raise DataError(f"Invalid sort entry at index {i}: column is required")
        if column not in known:
            raise DataError(f"Unknown sort column: {column!r}")
        if not isinstance(order_raw, str) or not order_raw:
            raise DataError(f"Invalid sort entry at index {i}: order is required")
        order = order_raw.strip().lower()
        if order not in _VALID_SORT_ORDERS:
            raise DataError(
                f"Invalid sort order: {order_raw!r}; expected 'asc' or 'desc'"
            )
        out.append({"column": column, "order": order})
    return {"by": out}


def _reject_cross_column_or(where: Sequence[Mapping[str, Any]]) -> None:
    """Refuse nested OR groups (use single-column ``in`` instead)."""
    for i, raw in enumerate(where):
        if not isinstance(raw, Mapping):
            continue
        op = str(raw.get("op", "")).strip().lower()
        if op != "or":
            continue
        nested = raw.get("clauses") or raw.get("where")
        cols = {
            str(c.get("column"))
            for c in (nested if isinstance(nested, (list, tuple)) else ())
            if isinstance(c, Mapping) and c.get("column")
        }
        if len(cols) > 1:
            raise DataError(
                "Cross-column OR is not viewable; use single-column 'in' or AND clauses"
            )
        raise DataError(
            "Nested OR groups are not supported; use operator 'in' on one column"
        )


def normalize_snapshot_pagination(
    limit: Any, offset: Any
) -> tuple[Optional[int], int]:
    """Parse optional limit/offset for snapshot (no default page size of 100).

    ``limit is None`` means write every remaining matching row after ``offset``.
    """
    if offset is None:
        off = 0
    else:
        try:
            off = int(offset)
        except (TypeError, ValueError) as e:
            raise DataError(f"Invalid offset: {offset!r}") from e
        if off < 0:
            raise DataError(f"Invalid offset: {offset!r}; must be >= 0")
    if limit is None:
        return None, off
    try:
        lim = int(limit)
    except (TypeError, ValueError) as e:
        raise DataError(f"Invalid limit: {limit!r}") from e
    if lim < 0:
        raise DataError(f"Invalid limit: {limit!r}; must be >= 0")
    return lim, off


def sort_table_rows(
    rows: Sequence[Mapping[str, Any]],
    sort: Optional[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Stable in-memory sort using a normalized ``view_spec.sort`` (``by`` list)."""
    out = [dict(r) for r in rows]
    if not sort:
        return out
    keys = sort.get("by") if isinstance(sort, Mapping) else None
    if not isinstance(keys, (list, tuple)) or not keys:
        return out

    def _part(value: Any) -> tuple[int, Any]:
        if value is None:
            return (2, "")
        if isinstance(value, bool):
            return (0, int(value))
        if isinstance(value, (int, float)):
            return (0, float(value))
        return (1, str(value).lower())

    for entry in reversed(list(keys)):
        col = str(entry["column"])
        reverse = str(entry.get("order", "asc")).lower() == "desc"
        out.sort(key=lambda row, c=col: _part(row.get(c)), reverse=reverse)
    return out


def build_snapshot_value_matrix(
    *,
    headers: Sequence[str],
    matching_rows: Sequence[Mapping[str, Any]],
    limit: Optional[int],
    offset: int,
) -> list[list[Any]]:
    """Header row plus projected values for a snapshot sheet (values only)."""
    if offset < 0:
        raise DataError(f"Invalid offset: {offset!r}; must be >= 0")
    if limit is None:
        page = list(matching_rows[offset:])
    else:
        if limit < 0:
            raise DataError(f"Invalid limit: {limit!r}; must be >= 0")
        page = list(matching_rows[offset : offset + limit])
    matrix: list[list[Any]] = [list(headers)]
    for row in page:
        matrix.append([row.get(h) for h in headers])
    return matrix


SNAPSHOT_SHEET_NAME_BASE = "mcp_view"


def allocate_unique_sheet_name(
    existing_names: Sequence[str],
    *,
    base: str = SNAPSHOT_SHEET_NAME_BASE,
) -> str:
    """First free name: ``base``, then ``base_2``, ``base_3``, ... (case-insensitive)."""
    taken = {str(n).strip().lower() for n in existing_names if str(n).strip()}
    candidate = base
    if candidate.lower() not in taken:
        return candidate
    n = 2
    while True:
        candidate = f"{base}_{n}"
        if candidate.lower() not in taken:
            return candidate
        n += 1


def validate_view_spec_for_apply(
    view_spec: Mapping[str, Any],
    *,
    mode: str,
    view_applicability: Optional[Mapping[str, Any]] = None,
    table_columns: Optional[Sequence[str]] = None,
) -> dict[str, Any]:
    """Validate ``view_spec`` + ``mode`` for ``apply_table_view``.

    Supports ``target.kind`` ``table`` (ListObject) and ``region`` (plain range).
    ``in_place`` refuses limit/offset and ``view_applicability.viewable is false``.
    ``snapshot`` honors limit/offset (no default page size) and does not treat
    limit/offset ``view_applicability`` as a blocker. Both modes refuse
    cross-column OR and non-viewable operators (e.g. ``is_empty``).
    """
    if not isinstance(view_spec, Mapping):
        raise DataError("view_spec must be an object")
    mode_norm = (mode or "").strip().lower()
    if mode_norm not in _VALID_VIEW_MODES:
        raise DataError(
            f"Invalid mode: {mode!r}; expected '{VIEW_MODE_IN_PLACE}' or "
            f"'{VIEW_MODE_SNAPSHOT}'"
        )

    snapshot_limit: Optional[int] = None
    snapshot_offset = 0
    if mode_norm == VIEW_MODE_IN_PLACE:
        if "limit" in view_spec and view_spec.get("limit") is not None:
            raise DataError(
                "view_spec must not include limit; Excel AutoFilter cannot show a row page"
            )
        if "offset" in view_spec and view_spec.get("offset"):
            raise DataError(
                "view_spec must not include offset; Excel AutoFilter cannot skip rows"
            )
        if view_applicability is not None:
            if not isinstance(view_applicability, Mapping):
                raise DataError("view_applicability must be an object when provided")
            if view_applicability.get("viewable") is False:
                reason = view_applicability.get("reason") or (
                    "view_applicability.viewable is false"
                )
                raise DataError(
                    f"Cannot apply view: {reason}. Sheet was not changed."
                )
    else:
        # Snapshot: limit/offset are part of the copied result set.
        if view_applicability is not None and not isinstance(
            view_applicability, Mapping
        ):
            raise DataError("view_applicability must be an object when provided")
        snapshot_limit, snapshot_offset = normalize_snapshot_pagination(
            view_spec.get("limit") if "limit" in view_spec else None,
            view_spec.get("offset") if "offset" in view_spec else None,
        )
        if isinstance(view_applicability, Mapping) and (
            view_applicability.get("viewable") is False
        ):
            reason = str(view_applicability.get("reason") or "").lower()
            page_reason = any(
                token in reason for token in ("limit", "offset", "truncat")
            )
            spec_has_page = (
                "limit" in view_spec and view_spec.get("limit") is not None
            ) or bool(view_spec.get("offset"))
            if page_reason and not spec_has_page:
                raise DataError(
                    "Cannot snapshot this query page: view_applicability says "
                    "rows were limited or offset, but view_spec has no limit or "
                    "offset. Add them so the sheet matches the queried rows. "
                    "Sheet was not changed."
                )

    target = view_spec.get("target")
    if not isinstance(target, Mapping):
        raise DataError("view_spec.target is required")
    kind = str(target.get("kind", "")).strip().lower()
    target_name: Optional[str] = None
    target_sheet: Optional[str] = None
    target_range: Optional[str] = None
    if kind == "table":
        name = target.get("name")
        if not isinstance(name, str) or not name.strip():
            raise DataError("view_spec.target.name is required for a table view")
        target_name = name.strip()
    elif kind == "region":
        sheet = target.get("sheet")
        range_raw = target.get("range")
        if not isinstance(sheet, str) or not sheet.strip():
            raise DataError("view_spec.target.sheet is required for a region view")
        if not isinstance(range_raw, str) or not range_raw.strip():
            raise DataError("view_spec.target.range is required for a region view")
        target_sheet = sheet.strip()
        target_range = normalize_excel_a1_range(range_raw)
    else:
        raise DataError(
            f"view_spec.target.kind {kind!r} is not supported for apply_table_view; "
            "expected 'table' or 'region'. Sheet was not changed."
        )

    columns_raw = view_spec.get("columns")
    if columns_raw is None:
        columns: list[str] = []
    elif isinstance(columns_raw, (list, tuple)):
        columns = [str(c) for c in columns_raw]
    else:
        raise DataError("view_spec.columns must be a list of column names")

    where_raw = view_spec.get("where")
    known = list(table_columns) if table_columns is not None else columns
    if table_columns is not None and columns:
        unknown = [c for c in columns if c not in set(table_columns)]
        if unknown:
            raise DataError(f"Unknown column: {unknown[0]!r}")
    # Detect nested OR before normalize_where_clauses rejects unknown ops.
    _reject_cross_column_or(where_raw if isinstance(where_raw, (list, tuple)) else [])
    clauses = normalize_where_clauses(
        where_raw if where_raw is not None else [],
        known_columns=known if known else columns,
    )

    for clause in clauses:
        if clause["op"] not in _VIEWABLE_OPS:
            raise DataError(
                f"Operator {clause['op']!r} is not viewable; sheet was not changed"
            )

    sort = normalize_view_sort(view_spec.get("sort"), known_columns=known if known else columns)
    out: dict[str, Any] = {
        "target_kind": kind,
        "columns": columns,
        "where": clauses,
        "sort": sort,
        "mode": mode_norm,
    }
    if kind == "table":
        out["target_name"] = target_name
    else:
        out["target_sheet"] = target_sheet
        out["target_range"] = target_range
    if mode_norm == VIEW_MODE_SNAPSHOT:
        out["limit"] = snapshot_limit
        out["offset"] = snapshot_offset
    return out


def _excel_filter_literal(value: Any) -> str:
    """Escape Excel AutoFilter wildcards so a contains pattern stays literal."""
    text = _as_text(value)
    return text.replace("~", "~~").replace("*", "~*").replace("?", "~?")


def criteria_text_for_clause(clause: Mapping[str, Any]) -> Any:
    """Map one viewable where clause to AutoFilter Criteria1 / operator payload."""
    op = clause["op"]
    value = clause.get("value")
    if op == "eq":
        return {"criteria1": value}
    if op == "neq":
        return {"criteria1": f"<>{_as_text(value)}"}
    if op == "contains":
        return {"criteria1": f"*{_excel_filter_literal(value)}*"}
    if op == "gt":
        return {"criteria1": f">{_as_text(value)}"}
    if op == "gte":
        return {"criteria1": f">={_as_text(value)}"}
    if op == "lt":
        return {"criteria1": f"<{_as_text(value)}"}
    if op == "lte":
        return {"criteria1": f"<={_as_text(value)}"}
    if op == "in":
        items = list(value) if isinstance(value, (list, tuple)) else [value]
        if len(items) == 0:
            raise DataError("Operator 'in' requires a non-empty list")
        if len(items) == 1:
            return {"criteria1": items[0]}
        if len(items) == 2:
            return {
                "criteria1": items[0],
                "criteria2": items[1],
                "operator": XL_FILTER_OR,
            }
        return {"criteria1": list(items), "operator": XL_FILTER_VALUES}
    raise DataError(f"Operator {op!r} is not viewable")


def compile_autofilter_field_steps(
    where: Sequence[Mapping[str, Any]],
    *,
    column_to_field: Mapping[str, int],
) -> list[dict[str, Any]]:
    """Compile AND-combined where clauses into per-field AutoFilter steps.

    ``column_to_field`` maps header name → 1-based field index relative to the
    table header (not worksheet column A).
    """
    # Multiple clauses on the same column are not expressible as one AutoFilter
    # field without AdvancedFilter; refuse rather than apply a partial view.
    seen: dict[str, Mapping[str, Any]] = {}
    for clause in where:
        col = str(clause["column"])
        if col in seen:
            raise DataError(
                f"Multiple where clauses on column {col!r} cannot be shown as one "
                "AutoFilter field; sheet was not changed"
            )
        seen[col] = clause

    steps: list[dict[str, Any]] = []
    for col, clause in seen.items():
        if col not in column_to_field:
            raise DataError(f"Unknown column: {col!r}")
        field = int(column_to_field[col])
        if field < 1:
            raise DataError(f"Invalid field index for column {col!r}")
        crit = criteria_text_for_clause(clause)
        step: dict[str, Any] = {"field": field, "column": col, **crit}
        steps.append(step)
    return steps


RESTORE_KIND_LISTOBJECT = "listobject"
RESTORE_KIND_RANGE = "range"
RESTORE_KIND_SNAPSHOT = "snapshot"
_VALID_RESTORE_KINDS = frozenset(
    {RESTORE_KIND_LISTOBJECT, RESTORE_KIND_RANGE, RESTORE_KIND_SNAPSHOT}
)


def build_restore_token(
    *,
    sheet: str,
    prior_filters: Sequence[Mapping[str, Any]] = (),
    prior_sort: Optional[Sequence[Mapping[str, Any]]] = None,
    sort_applied: bool = False,
    columns_hidden: Sequence[int] = (),
    kind: str = RESTORE_KIND_LISTOBJECT,
    table_name: Optional[str] = None,
    range_a1: Optional[str] = None,
    prior_autofilter_mode: Optional[bool] = None,
    created_by_tool: bool = False,
) -> dict[str, Any]:
    """Opaque-ish restore payload for ``clear_table_view`` (JSON-serializable)."""
    kind_norm = (kind or "").strip().lower()
    if kind_norm not in _VALID_RESTORE_KINDS:
        raise DataError(f"Unsupported restore_token kind: {kind!r}")
    if not isinstance(sheet, str) or not sheet.strip():
        raise DataError("sheet is required for restore_token")
    if kind_norm == RESTORE_KIND_SNAPSHOT:
        if not created_by_tool:
            raise DataError(
                "created_by_tool must be true for snapshot restore_token"
            )
        return {
            "v": RESTORE_TOKEN_VERSION,
            "kind": RESTORE_KIND_SNAPSHOT,
            "sheet": sheet.strip(),
            "created_by_tool": True,
        }
    token: dict[str, Any] = {
        "v": RESTORE_TOKEN_VERSION,
        "kind": kind_norm,
        "sheet": sheet.strip(),
        "prior_filters": [dict(f) for f in prior_filters],
        "prior_sort": [dict(s) for s in prior_sort] if prior_sort is not None else None,
        "sort_applied": bool(sort_applied),
        "columns_hidden": [int(c) for c in columns_hidden],
    }
    if kind_norm == RESTORE_KIND_LISTOBJECT:
        if not isinstance(table_name, str) or not table_name.strip():
            raise DataError("table_name is required for listobject restore_token")
        token["table"] = table_name.strip()
    else:
        if not isinstance(range_a1, str) or not range_a1.strip():
            raise DataError("range_a1 is required for range restore_token")
        if prior_autofilter_mode is None:
            raise DataError(
                "prior_autofilter_mode is required for range restore_token"
            )
        token["range"] = normalize_excel_a1_range(range_a1)
        token["prior_autofilter_mode"] = bool(prior_autofilter_mode)
    return token


def parse_restore_token(token: Mapping[str, Any]) -> dict[str, Any]:
    """Validate a restore_token from ``apply_table_view``."""
    if not isinstance(token, Mapping):
        raise DataError("restore_token must be an object")
    if token.get("v") != RESTORE_TOKEN_VERSION:
        raise DataError(f"Unsupported restore_token version: {token.get('v')!r}")
    kind = str(token.get("kind", "")).strip().lower()
    if kind not in _VALID_RESTORE_KINDS:
        raise DataError(
            f"restore_token.kind {token.get('kind')!r} is not supported here"
        )
    sheet = token.get("sheet")
    if not isinstance(sheet, str) or not sheet.strip():
        raise DataError("restore_token.sheet is required")
    if kind == RESTORE_KIND_SNAPSHOT:
        if token.get("created_by_tool") is not True:
            raise DataError(
                "restore_token.created_by_tool must be true to delete a snapshot sheet"
            )
        return {
            "kind": RESTORE_KIND_SNAPSHOT,
            "sheet": sheet.strip(),
            "created_by_tool": True,
        }
    prior_filters = token.get("prior_filters")
    if not isinstance(prior_filters, list):
        raise DataError("restore_token.prior_filters must be a list")
    columns_hidden = token.get("columns_hidden")
    if not isinstance(columns_hidden, list):
        raise DataError("restore_token.columns_hidden must be a list")
    base: dict[str, Any] = {
        "kind": kind,
        "sheet": sheet.strip(),
        "prior_filters": prior_filters,
        "prior_sort": token.get("prior_sort"),
        "sort_applied": bool(token.get("sort_applied")),
        "columns_hidden": [int(c) for c in columns_hidden],
    }
    if kind == RESTORE_KIND_LISTOBJECT:
        table = token.get("table")
        if not isinstance(table, str) or not table.strip():
            raise DataError("restore_token.table is required")
        base["table"] = table.strip()
        return base
    range_a1 = token.get("range")
    if not isinstance(range_a1, str) or not range_a1.strip():
        raise DataError("restore_token.range is required")
    if "prior_autofilter_mode" not in token:
        raise DataError("restore_token.prior_autofilter_mode is required")
    base["range"] = normalize_excel_a1_range(range_a1)
    base["prior_autofilter_mode"] = bool(token.get("prior_autofilter_mode"))
    return base


def build_query_table_payload(
    *,
    headers: Sequence[str],
    matching_rows: Sequence[Mapping[str, Any]],
    where: Sequence[Mapping[str, Any]],
    limit: int,
    offset: int,
    target: Mapping[str, Any],
) -> dict[str, Any]:
    """Page matching rows and attach ``view_spec`` / ``view_applicability``."""
    total = len(matching_rows)
    page = list(matching_rows[offset : offset + limit])
    truncated = (offset + len(page)) < total
    projected = [
        {h: row.get(h) for h in headers} for row in page
    ]
    # Keep every clause, including ones Excel cannot show. Epic-14 must refuse
    # a partial view when ``view_applicability.viewable`` is false.
    view_spec = {
        "target": dict(target),
        "columns": list(headers),
        "where": [dict(c) for c in where],
        "sort": None,
    }
    return {
        "headers": list(headers),
        "rows": projected,
        "row_count": len(projected),
        "truncated": truncated,
        "view_spec": view_spec,
        "view_applicability": build_view_applicability(
            where, offset=offset, truncated=truncated
        ),
    }


def filter_table_rows(
    rows: Sequence[Mapping[str, Any]],
    where: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Shared AND-combined predicate used by file and COM backends."""
    if not where:
        return [dict(r) for r in rows]
    out: list[dict[str, Any]] = []
    for row in rows:
        if all(evaluate_row_clause(row, clause) for clause in where):
            out.append(dict(row))
    return out


def query_table_rows(
    *,
    table_name: str,
    table_columns: Sequence[str],
    rows: Sequence[Mapping[str, Any]],
    columns: Optional[Sequence[str]] = None,
    where: Optional[Sequence[Mapping[str, Any]]] = None,
    limit: Optional[int] = None,
    offset: Optional[int] = None,
) -> dict[str, Any]:
    """Validate args, filter, page, and build the ``query_table`` payload."""
    headers = resolve_query_columns(table_columns, columns)
    clauses = normalize_where_clauses(where, known_columns=table_columns)
    lim, off = normalize_query_pagination(limit, offset)
    # Filter may reference columns outside the projection; rows must carry them.
    needed = set(headers) | {c["column"] for c in clauses}
    missing = needed - set(table_columns)
    if missing:
        raise DataError(f"Unknown column: {sorted(missing)[0]!r}")
    matching = filter_table_rows(rows, clauses)
    return build_query_table_payload(
        headers=headers,
        matching_rows=matching,
        where=clauses,
        limit=lim,
        offset=off,
        target={"kind": "table", "name": table_name},
    )


def query_region_rows(
    *,
    sheet: str,
    range_a1: str,
    region_columns: Sequence[str],
    rows: Sequence[Mapping[str, Any]],
    columns: Optional[Sequence[str]] = None,
    where: Optional[Sequence[Mapping[str, Any]]] = None,
    limit: Optional[int] = None,
    offset: Optional[int] = None,
) -> dict[str, Any]:
    """Validate args, filter, page, and build the ``query_region`` payload."""
    headers = resolve_query_columns(region_columns, columns)
    clauses = normalize_where_clauses(where, known_columns=region_columns)
    lim, off = normalize_query_pagination(limit, offset)
    needed = set(headers) | {c["column"] for c in clauses}
    missing = needed - set(region_columns)
    if missing:
        raise DataError(f"Unknown column: {sorted(missing)[0]!r}")
    matching = filter_table_rows(rows, clauses)
    return build_query_table_payload(
        headers=headers,
        matching_rows=matching,
        where=clauses,
        limit=lim,
        offset=off,
        target={
            "kind": "region",
            "sheet": sheet,
            "range": normalize_excel_a1_range(range_a1),
        },
    )


def _find_openpyxl_table(
    wb: Any, table_name: str
) -> tuple[Any, Any, str]:
    """Return ``(worksheet, table, display_name)`` or raise ``DataError``."""
    target = table_name.strip()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for key in list(ws.tables):
            table = ws.tables[key]
            display = str(table.displayName or table.name or key)
            if display == target or str(key) == target:
                return ws, table, display
    raise DataError(f"Table '{table_name}' not found")


def query_excel_table(
    filepath: str,
    table: str,
    columns: Optional[Sequence[str]] = None,
    where: Optional[Sequence[Mapping[str, Any]]] = None,
    limit: Optional[int] = None,
    offset: Optional[int] = None,
) -> dict[str, Any]:
    """Query a native Excel table via openpyxl without saving the workbook."""
    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        ws, tbl, display_name = _find_openpyxl_table(wb, table)
        table_columns = list(tbl.column_names)
        min_col, min_row, max_col, max_row = range_boundaries(
            normalize_excel_a1_range(str(tbl.ref))
        )
        headers_count = int(tbl.headerRowCount or 1) or 1
        totals = int(tbl.totalsRowCount or 0)
        data_start = min_row + headers_count
        data_end = max_row - totals
        rows: list[dict[str, Any]] = []
        if data_start <= data_end:
            for r in range(data_start, data_end + 1):
                row_map: dict[str, Any] = {}
                for i, col_name in enumerate(table_columns):
                    c = min_col + i
                    row_map[col_name] = ws.cell(row=r, column=c).value
                rows.append(row_map)
        return query_table_rows(
            table_name=display_name,
            table_columns=table_columns,
            rows=rows,
            columns=columns,
            where=where,
            limit=limit,
            offset=offset,
        )
    finally:
        wb.close()


def a1_range_from_bounds(min_col: int, min_row: int, max_col: int, max_row: int) -> str:
    """Format absolute 1-based bounds as an A1 range."""
    return (
        f"{get_column_letter(min_col)}{min_row}:"
        f"{get_column_letter(max_col)}{max_row}"
    )


def region_id_for(sheet: str, range_a1: str) -> str:
    """Stable region id: ``{sheet}!{range}``."""
    return f"{sheet}!{normalize_excel_a1_range(range_a1)}"


def parse_region_id(region_id: str, *, sheet_name: str) -> str:
    """Extract the A1 range from a region id; require matching sheet."""
    raw = (region_id or "").strip()
    if "!" not in raw:
        raise DataError(f"Invalid region_id: {region_id!r}")
    sheet, rng = raw.split("!", 1)
    if sheet != sheet_name:
        raise DataError(
            f"region_id sheet {sheet!r} does not match sheet_name {sheet_name!r}"
        )
    if not rng.strip():
        raise DataError(f"Invalid region_id: {region_id!r}")
    return normalize_excel_a1_range(rng)


def guess_header_from_row(values: Sequence[Any]) -> tuple[bool, list[str]]:
    """Return ``(header_guess, columns)``. Columns only when mostly text."""
    non_empty = [v for v in values if not _is_empty_cell(v)]
    if not non_empty:
        return False, []
    text_count = sum(1 for v in non_empty if _try_number(v) is None)
    if text_count <= len(non_empty) / 2:
        return False, []
    return True, [_as_text(v) if not _is_empty_cell(v) else "" for v in values]


def _island_overlaps_box(
    island: tuple[int, int, int, int],
    box: tuple[int, int, int, int],
) -> bool:
    """True when an island ``(min_row, min_col, max_row, max_col)`` hits a range box.

    ``box`` is openpyxl order: ``(min_col, min_row, max_col, max_row)``.
    """
    min_r, min_c, max_r, max_c = island
    b_min_c, b_min_r, b_max_c, b_max_r = box
    return not (
        max_r < b_min_r or min_r > b_max_r or max_c < b_min_c or min_c > b_max_c
    )


def _subtract_exclusion(
    island: tuple[int, int, int, int],
    box: tuple[int, int, int, int],
) -> list[tuple[int, int, int, int]]:
    """Remove one excluded box from an island, returning the leftover rectangles."""
    if not _island_overlaps_box(island, box):
        return [island]
    min_r, min_c, max_r, max_c = island
    b_min_c, b_min_r, b_max_c, b_max_r = box
    pieces: list[tuple[int, int, int, int]] = []
    if min_r < b_min_r:
        pieces.append((min_r, min_c, b_min_r - 1, max_c))
    if max_r > b_max_r:
        pieces.append((b_max_r + 1, min_c, max_r, max_c))
    mid_r1 = max(min_r, b_min_r)
    mid_r2 = min(max_r, b_max_r)
    if mid_r1 <= mid_r2 and min_c < b_min_c:
        pieces.append((mid_r1, min_c, mid_r2, b_min_c - 1))
    if mid_r1 <= mid_r2 and max_c > b_max_c:
        pieces.append((mid_r1, b_max_c + 1, mid_r2, max_c))
    return pieces


def _shrink_to_occupied(
    island: tuple[int, int, int, int],
    occupied: Sequence[Sequence[bool]],
    *,
    origin_row: int,
    origin_col: int,
) -> Optional[tuple[int, int, int, int]]:
    """Shrink a rectangle to occupied cells, or return None when none remain."""
    min_r, min_c, max_r, max_c = island
    found_r1 = found_c1 = found_r2 = found_c2 = None
    nrows = len(occupied)
    ncols = len(occupied[0]) if nrows else 0
    for r in range(min_r, max_r + 1):
        ir = r - origin_row
        if ir < 0 or ir >= nrows:
            continue
        for c in range(min_c, max_c + 1):
            ic = c - origin_col
            if ic < 0 or ic >= ncols or not occupied[ir][ic]:
                continue
            if found_r1 is None:
                found_r1 = found_r2 = r
                found_c1 = found_c2 = c
            else:
                found_r1 = min(found_r1, r)
                found_r2 = max(found_r2, r)
                found_c1 = min(found_c1, c)
                found_c2 = max(found_c2, c)
    if found_r1 is None:
        return None
    assert found_c1 is not None and found_r2 is not None and found_c2 is not None
    return (found_r1, found_c1, found_r2, found_c2)


def _islands_outside_exclusions(
    islands: Sequence[tuple[int, int, int, int]],
    boxes: Sequence[tuple[int, int, int, int]],
    occupied: Sequence[Sequence[bool]],
    *,
    origin_row: int,
    origin_col: int,
) -> list[tuple[int, int, int, int]]:
    """Drop ListObject cells from island bounds, then shrink to occupied cells."""
    out: list[tuple[int, int, int, int]] = []
    for island in islands:
        pieces = [island]
        for box in boxes:
            next_pieces: list[tuple[int, int, int, int]] = []
            for piece in pieces:
                next_pieces.extend(_subtract_exclusion(piece, box))
            pieces = next_pieces
        for piece in pieces:
            shrunk = _shrink_to_occupied(
                piece, occupied, origin_row=origin_row, origin_col=origin_col
            )
            if shrunk is not None:
                out.append(shrunk)
    return out


def _bounds_overlap_or_contain(
    r: int, c: int, boxes: Sequence[tuple[int, int, int, int]]
) -> bool:
    for min_col, min_row, max_col, max_row in boxes:
        if min_row <= r <= max_row and min_col <= c <= max_col:
            return True
    return False


def discover_occupied_islands(
    occupied: Sequence[Sequence[bool]],
    *,
    origin_row: int,
    origin_col: int,
) -> list[tuple[int, int, int, int]]:
    """Find rectangular islands separated by fully blank rows or columns.

    Returns absolute 1-based ``(min_row, min_col, max_row, max_col)`` tuples.
    """
    if not occupied:
        return []
    nrows = len(occupied)
    ncols = len(occupied[0]) if nrows else 0
    if nrows == 0 or ncols == 0:
        return []

    row_has = [any(occupied[r][c] for c in range(ncols)) for r in range(nrows)]
    bands: list[tuple[int, int]] = []
    r = 0
    while r < nrows:
        if not row_has[r]:
            r += 1
            continue
        r0 = r
        while r < nrows and row_has[r]:
            r += 1
        bands.append((r0, r - 1))

    islands: list[tuple[int, int, int, int]] = []
    for r0, r1 in bands:
        col_has = [
            any(occupied[rr][c] for rr in range(r0, r1 + 1)) for c in range(ncols)
        ]
        c = 0
        while c < ncols:
            if not col_has[c]:
                c += 1
                continue
            c0 = c
            while c < ncols and col_has[c]:
                c += 1
            c1 = c - 1
            min_r = max_r = min_c = max_c = None
            for rr in range(r0, r1 + 1):
                for cc in range(c0, c1 + 1):
                    if occupied[rr][cc]:
                        if min_r is None:
                            min_r = max_r = rr
                            min_c = max_c = cc
                        else:
                            min_r = min(min_r, rr)
                            max_r = max(max_r, rr)
                            min_c = min(min_c, cc)
                            max_c = max(max_c, cc)
            if min_r is not None:
                islands.append(
                    (
                        origin_row + min_r,
                        origin_col + min_c,
                        origin_row + max_r,
                        origin_col + max_c,
                    )
                )
    return islands


def build_region_layout_entry(
    *,
    sheet: str,
    min_row: int,
    min_col: int,
    max_row: int,
    max_col: int,
    header_values: Sequence[Any],
) -> dict[str, Any]:
    """One non-table island with a documented header guess."""
    range_a1 = a1_range_from_bounds(min_col, min_row, max_col, max_row)
    header_guess, columns = guess_header_from_row(header_values)
    header_range = a1_range_from_bounds(min_col, min_row, max_col, min_row)
    entry: dict[str, Any] = {
        "id": region_id_for(sheet, range_a1),
        "kind": "region",
        "sheet": sheet,
        "range": range_a1,
        "header_range": header_range,
        "header_row": min_row,
        "columns": columns,
        "header_guess": header_guess,
    }
    return entry


def build_sheet_layout_payload(
    tables: Sequence[Mapping[str, Any]],
    regions: Sequence[Mapping[str, Any]],
) -> dict[str, Any]:
    """Stable sheet-scoped layout envelope."""
    return {"tables": list(tables), "regions": list(regions)}


def layout_regions_from_grid(
    *,
    sheet: str,
    values: Sequence[Sequence[Any]],
    origin_row: int,
    origin_col: int,
    excluded_ranges: Sequence[str],
) -> list[dict[str, Any]]:
    """Discover islands outside ``excluded_ranges`` and attach header guesses."""
    if not values:
        return []
    nrows = len(values)
    ncols = max((len(row) for row in values), default=0)
    if nrows == 0 or ncols == 0:
        return []

    boxes: list[tuple[int, int, int, int]] = []
    for ref in excluded_ranges:
        try:
            boxes.append(range_boundaries(normalize_excel_a1_range(ref)))
        except Exception:
            continue

    occupied: list[list[bool]] = []
    for ir in range(nrows):
        row_vals = values[ir] if ir < len(values) else []
        row_flags: list[bool] = []
        for ic in range(ncols):
            abs_r = origin_row + ir
            abs_c = origin_col + ic
            cell = row_vals[ic] if ic < len(row_vals) else None
            in_table = _bounds_overlap_or_contain(abs_r, abs_c, boxes)
            row_flags.append(not in_table and not _is_empty_cell(cell))
        occupied.append(row_flags)

    islands = _islands_outside_exclusions(
        discover_occupied_islands(
            occupied, origin_row=origin_row, origin_col=origin_col
        ),
        boxes,
        occupied,
        origin_row=origin_row,
        origin_col=origin_col,
    )
    regions: list[dict[str, Any]] = []
    for min_r, min_c, max_r, max_c in islands:
        header_vals: list[Any] = []
        ir = min_r - origin_row
        for abs_c in range(min_c, max_c + 1):
            ic = abs_c - origin_col
            row_vals = values[ir] if 0 <= ir < len(values) else []
            header_vals.append(row_vals[ic] if 0 <= ic < len(row_vals) else None)
        regions.append(
            build_region_layout_entry(
                sheet=sheet,
                min_row=min_r,
                min_col=min_c,
                max_row=max_r,
                max_col=max_c,
                header_values=header_vals,
            )
        )
    return regions


def rows_from_header_and_data(
    headers: Sequence[str],
    data_matrix: Sequence[Sequence[Any]],
) -> list[dict[str, Any]]:
    """Zip data rows to header keys (pad/truncate to header width)."""
    out: list[dict[str, Any]] = []
    for raw in data_matrix:
        out.append({name: raw[i] if i < len(raw) else None for i, name in enumerate(headers)})
    return out


def resolve_region_query_source(
    *,
    matrix: Sequence[Sequence[Any]],
    range_a1: str,
    require_header_guess: bool,
) -> tuple[list[str], list[dict[str, Any]], str]:
    """Interpret first matrix row as headers (or require a text header guess)."""
    rng = normalize_excel_a1_range(range_a1)
    if not matrix:
        raise DataError(f"Region range {rng!r} is empty")
    header_row = list(matrix[0])
    if require_header_guess:
        ok, columns = guess_header_from_row(header_row)
        if not ok:
            raise DataError(
                "Region has no header guess; pass an explicit range that includes "
                "a header row"
            )
    else:
        columns = [
            _as_text(v) if not _is_empty_cell(v) else "" for v in header_row
        ]
    if not columns:
        raise DataError(f"Region range {rng!r} has no header row")
    data = [list(r) for r in matrix[1:]]
    rows = rows_from_header_and_data(columns, data)
    return columns, rows, rng


def _openpyxl_sheet_content_grid(
    ws: Any,
) -> tuple[Optional[tuple[int, int, int, int]], list[list[Any]]]:
    """Return content bounds and value grid, or ``(None, [])`` when empty."""
    max_row = int(ws.max_row or 1)
    max_col = int(ws.max_column or 1)
    min_r = min_c = max_r = max_c = None
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            if not _is_empty_cell(ws.cell(row=r, column=c).value):
                if min_r is None:
                    min_r = max_r = r
                    min_c = max_c = c
                else:
                    min_r = min(min_r, r)
                    max_r = max(max_r, r)
                    min_c = min(min_c, c)
                    max_c = max(max_c, c)
    if min_r is None:
        return None, []
    assert min_c is not None and max_r is not None and max_c is not None
    grid: list[list[Any]] = []
    for r in range(min_r, max_r + 1):
        grid.append(
            [ws.cell(row=r, column=c).value for c in range(min_c, max_c + 1)]
        )
    return (min_r, min_c, max_r, max_c), grid


def _openpyxl_sheet_tables(
    ws: Any, sheet_name: str, *, detail: str = DETAIL_SCHEMA
) -> list[dict[str, Any]]:
    """Catalog ListObjects on one sheet (layout ``kind: table`` entries)."""
    mode = validate_list_tables_detail(detail)
    tables: list[dict[str, Any]] = []
    for table_name in list(ws.tables):
        table = ws.tables[table_name]
        ref = str(table.ref)
        header_range, data_range, row_count = _table_ranges_from_ref(
            ref,
            header_row_count=int(table.headerRowCount or 1),
            totals_row_count=int(table.totalsRowCount or 0),
        )
        columns: Optional[list[str]] = None
        if mode == DETAIL_SCHEMA:
            columns = list(table.column_names)
        entry = build_table_catalog_entry(
            sheet=sheet_name,
            name=str(table.displayName or table.name or table_name),
            range_=ref,
            header_range=header_range,
            data_range=data_range,
            row_count=row_count,
            filter_applied=_openpyxl_filter_applied(table),
            columns=columns,
        )
        entry["kind"] = "table"
        tables.append(entry)
    return tables


def map_sheet_layout_from_worksheet(ws: Any, sheet_name: str) -> dict[str, Any]:
    """Map tables and islands for an already-open worksheet."""
    tables = _openpyxl_sheet_tables(ws, sheet_name)
    excluded = [str(t["range"]) for t in tables]
    bounds, grid = _openpyxl_sheet_content_grid(ws)
    if bounds is None:
        regions: list[dict[str, Any]] = []
    else:
        min_r, min_c, _, _ = bounds
        regions = layout_regions_from_grid(
            sheet=sheet_name,
            values=grid,
            origin_row=min_r,
            origin_col=min_c,
            excluded_ranges=excluded,
        )
    return build_sheet_layout_payload(tables, regions)


def map_excel_sheet_layout(filepath: str, sheet_name: str) -> dict[str, Any]:
    """Map native tables and non-table islands on one sheet (openpyxl, no save)."""
    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise DataError(f"Sheet '{sheet_name}' not found")
        return map_sheet_layout_from_worksheet(wb[sheet_name], sheet_name)
    finally:
        wb.close()


def _read_openpyxl_range_matrix(ws: Any, range_a1: str) -> list[list[Any]]:
    min_col, min_row, max_col, max_row = range_boundaries(
        normalize_excel_a1_range(range_a1)
    )
    matrix: list[list[Any]] = []
    for r in range(min_row, max_row + 1):
        matrix.append(
            [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
        )
    return matrix


def query_excel_region(
    filepath: str,
    sheet_name: str,
    *,
    region_id: Optional[str] = None,
    range: Optional[str] = None,  # noqa: A002 — MCP arg name
    columns: Optional[Sequence[str]] = None,
    where: Optional[Sequence[Mapping[str, Any]]] = None,
    limit: Optional[int] = None,
    offset: Optional[int] = None,
) -> dict[str, Any]:
    """Query a non-table region by id or explicit A1 range (openpyxl, no save)."""
    has_id = region_id is not None and str(region_id).strip() != ""
    has_range = range is not None and str(range).strip() != ""
    if has_id == has_range:
        raise DataError("Provide exactly one of region_id or range")

    wb = load_workbook(filepath, read_only=False, data_only=False)
    try:
        if sheet_name not in wb.sheetnames:
            raise DataError(f"Sheet '{sheet_name}' not found")
        ws = wb[sheet_name]
        if has_id:
            range_a1 = parse_region_id(str(region_id), sheet_name=sheet_name)
            layout = map_sheet_layout_from_worksheet(ws, sheet_name)
            known_ids = {str(r["id"]) for r in layout["regions"]}
            if str(region_id).strip() not in known_ids:
                raise DataError(f"Unknown region_id: {region_id!r}")
            require_guess = True
        else:
            range_a1 = normalize_excel_a1_range(str(range))
            require_guess = False

        matrix = _read_openpyxl_range_matrix(ws, range_a1)
        region_columns, rows, rng = resolve_region_query_source(
            matrix=matrix,
            range_a1=range_a1,
            require_header_guess=require_guess,
        )
        return query_region_rows(
            sheet=sheet_name,
            range_a1=rng,
            region_columns=region_columns,
            rows=rows,
            columns=columns,
            where=where,
            limit=limit,
            offset=offset,
        )
    finally:
        wb.close()


def create_excel_table(
    filepath: str,
    sheet_name: str,
    data_range: str,
    table_name: str | None = None,
    table_style: str = "TableStyleMedium9"
) -> dict:
    """Creates a native Excel table for the given data range.
    
    Args:
        filepath: Path to the Excel file.
        sheet_name: Name of the worksheet.
        data_range: The cell range for the table (e.g., "A1:D5").
        table_name: A unique name for the table. If not provided, a unique name is generated.
        table_style: The visual style to apply to the table.
        
    Returns:
        A dictionary with a success message and table details.
    """
    try:
        wb = load_workbook(filepath)
        if sheet_name not in wb.sheetnames:
            raise DataError(f"Sheet '{sheet_name}' not found.")
            
        ws = wb[sheet_name]

        # If no table name is provided, generate a unique one
        if not table_name:
            table_name = f"Table_{uuid.uuid4().hex[:8]}"

        # Check if table name already exists
        if table_name in ws.parent.defined_names:
            raise DataError(f"Table name '{table_name}' already exists.")

        # Create the table
        table = Table(displayName=table_name, ref=data_range)
        
        # Apply style
        style = TableStyleInfo(
            name=table_style, 
            showFirstColumn=False,
            showLastColumn=False, 
            showRowStripes=True, 
            showColumnStripes=False
        )
        table.tableStyleInfo = style
        
        ws.add_table(table)
        
        wb.save(filepath)
        
        return {
            "message": f"Successfully created table '{table_name}' in sheet '{sheet_name}'.",
            "table_name": table_name,
            "range": data_range
        }

    except Exception as e:
        logger.error(f"Failed to create table: {e}")
        raise DataError(str(e))
