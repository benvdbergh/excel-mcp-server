"""COM range read, export, metadata, validation, and direct-read fallback."""

from __future__ import annotations

import json
import os
from typing import Any, Callable, Dict, List, Optional, Tuple

from openpyxl.utils import get_column_letter

from excel_mcp.cells import parse_cell_range, validate_cell_reference
from excel_mcp.exceptions import DataError
from excel_mcp.query import (
    DEFAULT_EXPORT_MAX_ROWS,
    build_worksheet_table_payload,
    export_read_end_row,
    normalize_export_max_rows,
)
from excel_mcp.value_mode import validate_metadata_mode, validate_value_mode
from excel_mcp.com.session import (
    _coerce_com_count,
    get_open_workbook_com,
)

_XL_CELLTYPE_ALL_VALIDATION = 4

_XL_DVTYPE_STR: Dict[int, str] = {
    0: "none",
    1: "whole",
    2: "decimal",
    3: "list",
    4: "date",
    5: "time",
    6: "textLength",
    7: "custom",
}

def _is_excel_value_sequence(val: Any) -> bool:
    """True for list/tuple or SAFEARRAY-like sequences; never strings/bytes/mappings."""
    if val is None or isinstance(val, (str, bytes, bytearray, memoryview)):
        return False
    if isinstance(val, (dict, bool, int, float)):
        return False
    if isinstance(val, (list, tuple)):
        return True
    try:
        len(val)
        iter(val)
    except TypeError:
        return False
    return True


def _coerce_excel_cell_value(val: Any) -> Any:
    """Coerce a COM cell value to a JSON-friendly form.

    Sequences of scalars (tuple/list/SAFEARRAY-like) become nested lists of
    numbers/strings/bools/null so ``json.dumps(..., default=str)`` does not
    stringify them. Real strings are left unchanged.
    """
    if val is None or isinstance(val, (bool, int, float, str)):
        return val
    if not _is_excel_value_sequence(val):
        return val
    return [_coerce_excel_cell_value(item) for item in val]


def _normalize_excel_matrix(val: Any) -> List[List[Any]]:
    """Shape Excel ``Range.Value2`` / ``Formula`` into a list of rows (see pywin32 shapes)."""
    if val is None:
        return [[None]]
    if not isinstance(val, (tuple, list)):
        return [[_coerce_excel_cell_value(val)]]
    if len(val) == 0:
        return [[]]
    first = val[0]
    if isinstance(first, (tuple, list)) or _is_excel_value_sequence(first):
        return [
            [_coerce_excel_cell_value(c) for c in list(r)] for r in val
        ]
    return [[_coerce_excel_cell_value(c) for c in val]]

def _is_blankish_excel_value(val: Any) -> bool:
    """Treat Excel-empty values as blank for sparse matrix heuristics."""
    return val is None or val == ""


def _matrix_is_entirely_blankish(
    matrix: List[List[Any]], nrows: int, ncols: int
) -> bool:
    """True when every cell in the bulk matrix looks blank for the given shape."""
    for ir in range(nrows):
        row_vals = matrix[ir] if ir < len(matrix) else []
        for ic in range(ncols):
            bulk_val = row_vals[ic] if ic < len(row_vals) else None
            if not _is_blankish_excel_value(bulk_val):
                return False
    return True

def _direct_read_matrix_com(
    ws: Any, srow: int, scol: int, nrows: int, ncols: int
) -> List[List[Any]]:
    """Read a rectangular range cell-by-cell via ``Worksheet.Cells``."""
    out: List[List[Any]] = []
    for ir in range(nrows):
        r = srow + ir
        row_vals: List[Any] = []
        for ic in range(ncols):
            c = scol + ic
            try:
                row_vals.append(_coerce_excel_cell_value(ws.Cells(r, c).Value2))
            except Exception:
                row_vals.append(None)
        out.append(row_vals)
    return out

def _direct_read_matrix_text_com(
    ws: Any, srow: int, scol: int, nrows: int, ncols: int
) -> List[List[Any]]:
    """Read displayed text cell-by-cell via ``Worksheet.Cells`` ``Text``."""
    out: List[List[Any]] = []
    for ir in range(nrows):
        r = srow + ir
        row_vals: List[Any] = []
        for ic in range(ncols):
            c = scol + ic
            try:
                row_vals.append(_coerce_excel_cell_value(ws.Cells(r, c).Text))
            except Exception:
                row_vals.append(None)
        out.append(row_vals)
    return out

_FALLBACK_LARGE_RANGE_CELLS = 64

_FALLBACK_MAX_CHECKS = 24

_FALLBACK_MISMATCH_THRESHOLD = 3

def _evenly_spaced_indices(count: int, samples: int) -> List[int]:
    """Return ``samples`` indices in ``[0, count)``, evenly spaced incl. endpoints."""
    if count <= 0:
        return []
    if samples >= count:
        return list(range(count))
    if samples <= 1:
        return [0]
    step = (count - 1) / (samples - 1)
    return [round(i * step) for i in range(samples)]

def _fallback_sample_coords(
    nrows: int, ncols: int, max_coords: int, *, stratified: bool
) -> List[Tuple[int, int]]:
    """Coordinates to probe for bulk-vs-direct sparsity anomalies."""
    if not stratified:
        return [(ir, ic) for ir in range(nrows) for ic in range(ncols)]
    n_row_samples = max(1, min(nrows, int(max_coords**0.5) + 1))
    n_col_samples = max(1, min(ncols, max(1, max_coords // n_row_samples)))
    row_indices = _evenly_spaced_indices(nrows, n_row_samples)
    col_indices = _evenly_spaced_indices(ncols, n_col_samples)
    coords: List[Tuple[int, int]] = []
    for ir in row_indices:
        for ic in col_indices:
            coords.append((ir, ic))
            if len(coords) >= max_coords:
                return coords
    return coords

def _should_fallback_to_direct_read_impl(
    ws: Any,
    matrix: List[List[Any]],
    srow: int,
    scol: int,
    nrows: int,
    ncols: int,
    read_cell: Callable[[int, int], Any],
) -> bool:
    """Probe bulk matrix vs direct cell reads; fallback when mismatches accumulate.

    Sampling: up to ``_FALLBACK_MAX_CHECKS`` blank-looking bulk cells; trigger when
    mismatches reach the threshold. When the bulk matrix is entirely blankish, a
    single direct non-blank is enough (sparse wide ranges can hide a handful of
    values). Otherwise keep ``_FALLBACK_MISMATCH_THRESHOLD`` so slightly sparse
    sheets stay on the bulk path. Small ranges
    (``rows * cols <= _FALLBACK_LARGE_RANGE_CELLS``) scan row-major from the top-left;
    larger ranges use a stratified grid so late rows/columns are not under-sampled.
    """
    if nrows <= 0 or ncols <= 0:
        return False

    mismatch_threshold = (
        1
        if _matrix_is_entirely_blankish(matrix, nrows, ncols)
        else _FALLBACK_MISMATCH_THRESHOLD
    )
    stratified = (nrows * ncols) > _FALLBACK_LARGE_RANGE_CELLS
    coords = _fallback_sample_coords(
        nrows, ncols, _FALLBACK_MAX_CHECKS, stratified=stratified
    )
    checked = 0
    mismatches = 0
    for ir, ic in coords:
        if checked >= _FALLBACK_MAX_CHECKS:
            break
        row_vals = matrix[ir] if ir < len(matrix) else []
        bulk_val = row_vals[ic] if ic < len(row_vals) else None
        if not _is_blankish_excel_value(bulk_val):
            continue
        r = srow + ir
        c = scol + ic
        try:
            direct_val = read_cell(r, c)
        except Exception:
            direct_val = None
        checked += 1
        if not _is_blankish_excel_value(direct_val):
            mismatches += 1
            if mismatches >= mismatch_threshold:
                return True
    return False

def _should_fallback_to_direct_read_com(
    ws: Any,
    matrix: List[List[Any]],
    srow: int,
    scol: int,
    nrows: int,
    ncols: int,
) -> bool:
    """Detect COM ``Range.Value2`` sparsity anomalies and trigger direct cell reads."""
    return _should_fallback_to_direct_read_impl(
        ws,
        matrix,
        srow,
        scol,
        nrows,
        ncols,
        lambda r, c: ws.Cells(r, c).Value2,
    )

def _should_fallback_to_direct_read_text_com(
    ws: Any,
    matrix: List[List[Any]],
    srow: int,
    scol: int,
    nrows: int,
    ncols: int,
) -> bool:
    """Detect bulk ``Range.Text`` sparsity anomalies (mirrors Value2 resiliency)."""
    return _should_fallback_to_direct_read_impl(
        ws,
        matrix,
        srow,
        scol,
        nrows,
        ncols,
        lambda r, c: ws.Cells(r, c).Text,
    )

def _read_range_matrix_com(
    ws: Any,
    rng: Any,
    srow: int,
    scol: int,
    nrows: int,
    ncols: int,
    *,
    value_mode: str = "value",
) -> List[List[Any]]:
    """Read a rectangular COM range as a matrix (Value2 or displayed Text)."""
    if value_mode == "text":
        if nrows == 1 and ncols == 1:
            try:
                matrix = _normalize_excel_matrix(rng.Text)
            except Exception:
                return _direct_read_matrix_text_com(ws, srow, scol, nrows, ncols)
            bulk_val = matrix[0][0] if matrix and matrix[0] else None
            if _is_blankish_excel_value(bulk_val):
                direct = _direct_read_matrix_text_com(ws, srow, scol, nrows, ncols)
                direct_val = direct[0][0] if direct and direct[0] else None
                if not _is_blankish_excel_value(direct_val):
                    return direct
            elif _should_fallback_to_direct_read_text_com(
                ws, matrix, srow, scol, nrows, ncols
            ):
                return _direct_read_matrix_text_com(ws, srow, scol, nrows, ncols)
            return matrix
        return _direct_read_matrix_text_com(ws, srow, scol, nrows, ncols)

    try:
        raw_vals = rng.Value2
    except Exception:
        return _direct_read_matrix_com(ws, srow, scol, nrows, ncols)
    matrix = _normalize_excel_matrix(raw_vals)
    if _should_fallback_to_direct_read_com(ws, matrix, srow, scol, nrows, ncols):
        return _direct_read_matrix_com(ws, srow, scol, nrows, ncols)
    return matrix

def _com_used_bounds(ws: Any) -> Tuple[int, int, int, int]:
    """Return min_row, min_col, max_row, max_col from ``Worksheet.UsedRange``."""
    try:
        u = ws.UsedRange
        r0 = int(u.Row)
        c0 = int(u.Column)
        nr = int(u.Rows.Count)
        nc = int(u.Columns.Count)
        return r0, c0, r0 + nr - 1, c0 + nc - 1
    except Exception:
        return 1, 1, 1, 1

def _com_validation_dict(ws: Any, row: int, col: int, cell_address: str) -> Dict[str, Any]:
    """Map COM ``Validation`` to the openpyxl-oriented dict used in ``cell_validation``."""
    try:
        cell = ws.Cells(row, col)
        v = cell.Validation
        vtype = int(v.Type)
    except Exception:
        return {"has_validation": False}

    if vtype == 0:
        return {"has_validation": False}

    type_str = _XL_DVTYPE_STR.get(vtype, "custom")
    info: Dict[str, Any] = {
        "cell": cell_address,
        "has_validation": True,
        "validation_type": type_str,
    }
    try:
        info["allow_blank"] = bool(getattr(v, "IgnoreBlank", True))
    except Exception:
        info["allow_blank"] = True
    try:
        op = getattr(v, "Operator", None)
        if op is not None and int(op) != 0:
            info["operator"] = int(op)
    except Exception:
        pass
    for attr, key in (
        ("InputMessage", "prompt"),
        ("InputTitle", "prompt_title"),
        ("ErrorMessage", "error_message"),
        ("ErrorTitle", "error_title"),
    ):
        try:
            s = str(getattr(v, attr, "") or "").strip()
            if s:
                info[key] = s
        except Exception:
            continue
    try:
        f1 = str(getattr(v, "Formula1", "") or "").lstrip("=")
        f2 = str(getattr(v, "Formula2", "") or "").lstrip("=")
    except Exception:
        f1, f2 = "", ""
    if type_str == "list" and f1:
        if "," in f1 and "!" not in f1:
            info["allowed_values"] = [x.strip() for x in f1.split(",") if x.strip()]
        else:
            info["allowed_values"] = [f1] if f1 else []
            if f1 and ("," not in f1) and f1 not in info["allowed_values"]:
                info["allowed_values"] = [f1]
    else:
        if f1:
            info["formula1"] = f1
        if f2:
            info["formula2"] = f2
    return info

def _com_workbook_metadata_dict(wb: Any, filepath: str, include_ranges: bool) -> Dict[str, Any]:
    """Build a JSON-like dict aligned with :func:`excel_mcp.workbook.get_workbook_info`."""
    from pathlib import Path

    path = Path(filepath)
    names: list[str] = []
    try:
        n = _coerce_com_count(getattr(wb.Worksheets, "Count", 0))
    except Exception:
        n = 0
    for i in range(1, n + 1):
        try:
            names.append(str(wb.Worksheets.Item(i).Name))
        except Exception:
            continue

    fn = filepath
    try:
        fn = str(wb.FullName)
    except Exception:
        pass
    filename = os.path.basename(fn.rstrip("/")) or fn

    info: Dict[str, Any] = {
        "filename": filename,
        "sheets": names,
        "size": 0,
        "modified": 0.0,
    }
    disk_path: Optional[str] = None
    try:
        p = str(getattr(wb, "Path", "") or "").strip()
        nm = str(getattr(wb, "Name", "") or "").strip()
        if p and nm:
            disk_path = os.path.normpath(os.path.join(p, nm))
    except Exception:
        disk_path = None
    if disk_path and os.path.isfile(disk_path):
        try:
            st = os.stat(disk_path)
            info["size"] = st.st_size
            info["modified"] = st.st_mtime
        except OSError:
            pass
    elif path.is_file():
        try:
            st = path.stat()
            info["size"] = st.st_size
            info["modified"] = st.st_mtime
        except OSError:
            pass

    if include_ranges:
        ranges: Dict[str, str] = {}
        for i in range(1, n + 1):
            try:
                ws = wb.Worksheets.Item(i)
                name = str(ws.Name)
                mr, mc = _com_used_bounds(ws)[2:]
                if mr >= 1 and mc >= 1:
                    ranges[name] = f"A1:{get_column_letter(mc)}{mr}"
            except Exception:
                continue
        info["used_ranges"] = ranges
    return info

def _com_sheet_merge_addresses(ws: Any) -> List[str]:
    """Return merged range addresses (e.g. ``A1:B2``) for a COM worksheet."""
    seen: set[str] = set()
    out: List[str] = []
    try:
        u = ws.UsedRange
    except Exception:
        return []
    try:
        r0, c0 = int(u.Row), int(u.Column)
        nr, nc = int(u.Rows.Count), int(u.Columns.Count)
    except Exception:
        return []
    for r in range(r0, r0 + nr):
        for c in range(c0, c0 + nc):
            try:
                cell = ws.Cells(r, c)
                if _com_bool_is_true(getattr(cell, "MergeCells", False)):
                    area = cell.MergeArea
                    addr = str(getattr(area, "Address", "")).replace("$", "")
                    if addr and addr not in seen:
                        seen.add(addr)
                        out.append(addr)
            except Exception:
                continue
    return out

def _com_validation_rules_for_sheet(ws: Any) -> List[Dict[str, Any]]:
    """Build a list similar to :func:`excel_mcp.cell_validation.get_all_validation_ranges`."""
    out: List[Dict[str, Any]] = []
    try:
        u = ws.UsedRange
    except Exception:
        return []
    try:
        spec = u.SpecialCells(_XL_CELLTYPE_ALL_VALIDATION)
    except Exception:
        return []
    try:
        n_areas = int(spec.Areas.Count)
    except Exception:
        n_areas = 1
    for i in range(1, n_areas + 1):
        try:
            area = spec.Areas(i) if n_areas > 1 else spec
            addr = str(area.Address).replace("$", "")
            v = area.Cells(1, 1).Validation
            vt = int(getattr(v, "Type", 0))
            if vt == 0:
                continue
            type_str = _XL_DVTYPE_STR.get(vt, "custom")
            info: Dict[str, Any] = {
                "ranges": addr,
                "validation_type": type_str,
                "allow_blank": bool(getattr(v, "IgnoreBlank", True)),
            }
            if type_str == "list":
                f1 = str(getattr(v, "Formula1", "") or "").lstrip("=")
                if f1 and "," in f1 and "!" not in f1:
                    info["allowed_values"] = [x.strip() for x in f1.split(",") if x.strip()]
                elif f1:
                    info["allowed_values"] = [f1]
            else:
                f1 = str(getattr(v, "Formula1", "") or "").lstrip("=")
                f2 = str(getattr(v, "Formula2", "") or "").lstrip("=")
                if f1:
                    info["formula1"] = f1
                if f2:
                    info["formula2"] = f2
            out.append(info)
        except Exception:
            continue
    return out

def read_range_with_metadata_com(
    filepath: str,
    sheet_name: str,
    start_cell: str,
    end_cell: Optional[str],
    value_mode: str = "value",
    metadata_mode: str = "full",
) -> str:
    include_validation = metadata_mode == "full"
    raw_start = start_cell.strip()
    ec: Optional[str] = end_cell
    if ":" in raw_start and ec is None:
        parts = raw_start.split(":", 1)
        raw_start, ec = parts[0].strip(), parts[1].strip()

    wb_com, err = get_open_workbook_com(filepath)
    if err:
        return err
    try:
        ws = wb_com.Worksheets(sheet_name)
    except Exception:
        return f"Error: Sheet '{sheet_name}' not found"

    try:
        srow, scol, _, _ = parse_cell_range(f"{raw_start}:{raw_start}")
    except ValueError as e:
        return f"Error: Invalid start cell format: {str(e)}"

    uminr, uminc, umaxr, umaxc = _com_used_bounds(ws)
    if ec is None:
        if umaxr == 1 and umaxc == 1:
            try:
                a1v = ws.Range("A1").Value2
            except Exception:
                a1v = None
            if a1v is None or a1v == "":
                erow, ecol = srow, scol
            else:
                erow, ecol = umaxr, umaxc
                if raw_start.upper() == "A1":
                    srow, scol = uminr, uminc
        else:
            erow, ecol = umaxr, umaxc
            if raw_start.upper() == "A1":
                srow, scol = uminr, uminc
    else:
        try:
            et = ec.strip() if ec else ""
            erow, ecol, _, _ = parse_cell_range(f"{et}:{et}")
        except ValueError as e:
            return f"Error: Invalid end cell format: {str(e)}"

    if srow > umaxr or scol > umaxc:
        return json.dumps(
            {
                "range": f"{get_column_letter(scol)}{srow}:",
                "sheet_name": sheet_name,
                "value_mode": value_mode,
                "metadata_mode": metadata_mode,
                "cells": [],
            },
            indent=2,
            default=str,
        )

    range_str = (
        f"{get_column_letter(scol)}{srow}:{get_column_letter(ecol)}{erow}"
    )
    range_data: Dict[str, Any] = {
        "range": range_str,
        "sheet_name": sheet_name,
        "value_mode": value_mode,
        "metadata_mode": metadata_mode,
        "cells": [],
    }

    nrows = erow - srow + 1
    ncols = ecol - scol + 1
    try:
        top_left = ws.Cells(srow, scol)
        rng = top_left.Resize(nrows, ncols)
        matrix = _read_range_matrix_com(
            ws, rng, srow, scol, nrows, ncols, value_mode=value_mode
        )
    except Exception as exc:
        return f"Error: {exc}"
    for ir in range(nrows):
        row_vals = matrix[ir] if ir < len(matrix) else []
        for ic in range(ncols):
            r = srow + ir
            c = scol + ic
            val: Any
            if ir < len(matrix) and ic < len(row_vals):
                val = row_vals[ic]
            else:
                val = None
            addr = f"{get_column_letter(c)}{r}"
            cell_data: Dict[str, Any] = {
                "address": addr,
                "value": val,
                "row": r,
                "column": c,
            }
            if include_validation:
                vinfo = _com_validation_dict(ws, r, c, addr)
                if vinfo:
                    cell_data["validation"] = vinfo
                else:
                    cell_data["validation"] = {"has_validation": False}
            range_data["cells"].append(cell_data)

    if not range_data["cells"]:
        return "No data found in specified range"
    return json.dumps(range_data, indent=2, default=str)

def export_worksheet_table_com(
    filepath: str,
    sheet_name: str,
    start_cell: str,
    end_cell: Optional[str],
    max_rows: int,
) -> str:
    try:
        cap = normalize_export_max_rows(max_rows)
    except DataError as e:
        return f"Error: {str(e)}"

    raw_start = start_cell.strip()
    ec: Optional[str] = end_cell
    if ":" in raw_start and ec is None:
        parts = raw_start.split(":", 1)
        raw_start, ec = parts[0].strip(), parts[1].strip()

    wb_com, err = get_open_workbook_com(filepath)
    if err:
        return err
    try:
        ws = wb_com.Worksheets(sheet_name)
    except Exception:
        return f"Error: Sheet '{sheet_name}' not found"

    try:
        srow, scol, _, _ = parse_cell_range(f"{raw_start}:{raw_start}")
    except ValueError as e:
        return f"Error: Invalid start cell format: {str(e)}"

    uminr, uminc, umaxr, umaxc = _com_used_bounds(ws)
    if ec is None:
        if umaxr == 1 and umaxc == 1:
            try:
                a1v = ws.Range("A1").Value2
            except Exception:
                a1v = None
            if a1v is None or a1v == "":
                erow, ecol = srow, scol
            else:
                erow, ecol = umaxr, umaxc
                if raw_start.upper() == "A1":
                    srow, scol = uminr, uminc
        else:
            erow, ecol = umaxr, umaxc
            if raw_start.upper() == "A1":
                srow, scol = uminr, uminc
    else:
        try:
            et = ec.strip() if ec else ""
            erow, ecol, _, _ = parse_cell_range(f"{et}:{et}")
        except ValueError as e:
            return f"Error: Invalid end cell format: {str(e)}"

    if srow > umaxr or scol > umaxc:
        range_str = f"{get_column_letter(scol)}{srow}:"
        payload = build_worksheet_table_payload(
            sheet_name, range_str, [], max_rows=cap
        )
        return json.dumps(payload, indent=2, default=str)

    range_str = (
        f"{get_column_letter(scol)}{srow}:{get_column_letter(ecol)}{erow}"
    )
    read_erow = export_read_end_row(srow, erow, cap)
    total_data_rows = max(0, erow - srow)
    nrows = read_erow - srow + 1
    ncols = ecol - scol + 1
    try:
        top_left = ws.Cells(srow, scol)
        rng = top_left.Resize(nrows, ncols)
        matrix = _read_range_matrix_com(
            ws, rng, srow, scol, nrows, ncols, value_mode="value"
        )
    except Exception as exc:
        return f"Error: {exc}"

    payload = build_worksheet_table_payload(
        sheet_name,
        range_str,
        matrix,
        max_rows=cap,
        total_data_rows=total_data_rows,
    )
    return json.dumps(payload, indent=2, default=str)

def workbook_metadata_com(filepath: str, include_ranges: bool) -> str:
    wb_com, err = get_open_workbook_com(filepath)
    if err:
        return err
    info = _com_workbook_metadata_dict(wb_com, filepath, include_ranges)
    return str(info)

def read_merged_cell_ranges_com(filepath: str, sheet_name: str) -> str:
    wb_com, err = get_open_workbook_com(filepath)
    if err:
        return err
    try:
        ws = wb_com.Worksheets(sheet_name)
    except Exception:
        return f"Error: Sheet '{sheet_name}' not found"
    return str(_com_sheet_merge_addresses(ws))

def read_worksheet_data_validation_com(
    filepath: str, sheet_name: str
) -> str:
    wb_com, err = get_open_workbook_com(filepath)
    if err:
        return err
    try:
        ws = wb_com.Worksheets(sheet_name)
    except Exception:
        return f"Error: Sheet '{sheet_name}' not found"
    rules = _com_validation_rules_for_sheet(ws)
    if not rules:
        return "No data validation rules found in this worksheet"
    return json.dumps(
        {"sheet_name": sheet_name, "validation_rules": rules},
        indent=2,
        default=str,
    )

def validate_sheet_range_com(
    filepath: str,
    sheet_name: str,
    start_cell: str,
    end_cell: Optional[str],
) -> str:
    from openpyxl.utils import get_column_letter as gcl

    wb_com, err = get_open_workbook_com(filepath)
    if err:
        return err
    try:
        ws = wb_com.Worksheets(sheet_name)
    except Exception:
        return f"Error: Sheet '{sheet_name}' not found"

    try:
        srow, scol, erow, ecol = parse_cell_range(start_cell, end_cell)
    except ValueError as e:
        return f"Error: Invalid range: {str(e)}"
    if erow is None:
        erow = srow
    if ecol is None:
        ecol = scol

    _uminr, _uminc, umaxr, umaxc = _com_used_bounds(ws)
    data_max_row, data_max_col = umaxr, umaxc
    data_range_str = f"A1:{gcl(data_max_col)}{data_max_row}"
    range_str = start_cell if end_cell is None else f"{start_cell}:{end_cell}"
    return (
        f"Range '{range_str}' is valid. "
        f"Sheet contains data in range '{data_range_str}'"
    )

def validate_formula_syntax_com(
    filepath: str, sheet_name: str, cell: str, formula: str
) -> str:
    from excel_mcp.formula_syntax import validate_formula

    if not validate_cell_reference(cell):
        return f"Error: Invalid cell reference: {cell}"
    ftext = formula if formula.startswith("=") else f"={formula}"
    is_valid, message = validate_formula(ftext)
    if not is_valid:
        return f"Error: Invalid formula syntax: {message}"

    wb_com, err = get_open_workbook_com(filepath)
    if err:
        return err
    try:
        ws = wb_com.Worksheets(sheet_name)
    except Exception:
        return f"Error: Sheet '{sheet_name}' not found"

    try:
        current = ws.Range(cell).Formula
    except Exception as exc:
        return f"Error: {exc}"
    if isinstance(current, str) and current.startswith("="):
        if ftext == current:
            return "Formula is valid and matches cell content"
        return (
            "Formula is valid but doesn't match cell content; "
            f"cell has {current!r}, provided {ftext!r}"
        )
    return (
        "Formula is valid but cell contains no formula; "
        f"current content: {current!r}"
    )

def evaluate_range_com(
    filepath: str,
    sheet_name: str,
    start_cell: Optional[str],
    end_cell: Optional[str],
) -> str:
    from excel_mcp.cells import validate_cell_reference

    if end_cell is not None and start_cell is None:
        return "Error: end_cell requires start_cell"
    if start_cell is not None and not validate_cell_reference(start_cell):
        return f"Error: Invalid start cell reference: {start_cell}"
    if end_cell is not None and not validate_cell_reference(end_cell):
        return f"Error: Invalid end cell reference: {end_cell}"

    wb_com, err = get_open_workbook_com(filepath)
    if err:
        return err
    try:
        ws = wb_com.Worksheets(sheet_name)
    except Exception:
        return f"Error: Sheet '{sheet_name}' not found"

    try:
        if start_cell is None:
            ws.Calculate()
            scope = f"sheet '{sheet_name}'"
        elif end_cell is None:
            ws.Range(start_cell).Calculate()
            scope = f"range {start_cell}"
        else:
            ws.Range(start_cell, end_cell).Calculate()
            scope = f"range {start_cell}:{end_cell}"
    except Exception as exc:
        return f"Error: {exc}"

    return (
        f"Recalculated {scope} via Excel COM (in-memory only; "
        f"call save_workbook to flush to disk for file-based reads)"
    )

